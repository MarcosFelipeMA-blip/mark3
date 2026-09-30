"""Estilos e utilidades comuns das planilhas."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule

BRL = 'R$ #,##0.00'
PCT = '0.0%'
DATA = 'DD/MM/YYYY'
INPUT = "FFF4CC"
thin = Side(style="thin", color="D9D9D9")
BORDA = Border(left=thin, right=thin, top=thin, bottom=thin)


class Tema:
    def __init__(self, cor, suave):
        self.cor, self.suave = cor, suave
        self.H = Font(bold=True, color="FFFFFF")
        self.HF = PatternFill("solid", fgColor=cor)
        self.SF = PatternFill("solid", fgColor=suave)
        self.IF = PatternFill("solid", fgColor=INPUT)

    def titulo(self, ws, texto, sub=None):
        ws["A1"] = texto
        ws["A1"].font = Font(bold=True, size=16, color=self.cor)
        if sub:
            ws["A2"] = sub
            ws["A2"].font = Font(italic=True, color="666666")

    def cabecalho(self, ws, linha, colunas, altura=32):
        for i, c in enumerate(colunas, 1):
            cell = ws.cell(row=linha, column=i, value=c)
            cell.font, cell.fill, cell.border = self.H, self.HF, BORDA
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[linha].height = altura

    def entrada(self, cell, fmt=None):
        cell.fill, cell.border = self.IF, BORDA
        if fmt:
            cell.number_format = fmt

    def calculo(self, cell, fmt=None, negrito=False):
        cell.fill, cell.border = self.SF, BORDA
        if fmt:
            cell.number_format = fmt
        if negrito:
            cell.font = Font(bold=True)

    def rotulo(self, ws, ref, texto, negrito=True):
        ws[ref] = texto
        ws[ref].font = Font(bold=negrito)

    def nota(self, ws, ref, texto):
        ws[ref] = texto
        ws[ref].font = Font(italic=True, color="666666")


def larguras(ws, valores):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(valores, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def vermelho_se(ws, faixa, formula):
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[formula],
                                  fill=PatternFill("solid", fgColor="FDE2E1"), font=Font(color="A12622", bold=True)))


def verde_se(ws, faixa, formula):
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[formula],
                                  fill=PatternFill("solid", fgColor="E3F4E8"), font=Font(color="1E6B3C", bold=True)))


def amarelo_se(ws, faixa, formula):
    ws.conditional_formatting.add(faixa, FormulaRule(formula=[formula],
                                  fill=PatternFill("solid", fgColor="FFF1C2"), font=Font(color="8A5A00", bold=True)))


def como_usar(ws, tema, titulo, passos):
    tema.titulo(ws, titulo, "Preencha só as células AMARELAS. As células coloridas são calculadas sozinhas.")
    for i, s in enumerate(passos, 4):
        ws.cell(row=i, column=1, value=s).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
