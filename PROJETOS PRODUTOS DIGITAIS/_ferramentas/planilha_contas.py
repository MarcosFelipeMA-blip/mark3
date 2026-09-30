import sys
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment
from estilo import Tema, BRL, PCT, DATA, larguras, vermelho_se, verde_se, amarelo_se, como_usar

t = Tema("1E7A4C", "EDF7F1")
wb = Workbook()

como_usar(wb.active, t, "Planilhas Contas em Dia", [
    "COMO USAR (na ordem)",
    "1. 'Orçamento': coloque o mês de referência (célula C2), suas entradas e quanto PLANEJA gastar em cada categoria.",
    "2. 'Controle de Gastos': registre cada gasto do dia a dia (2 minutos por dia). O 'Real' do Orçamento é preenchido sozinho.",
    "3. 'Minhas Dívidas': liste TODAS as dívidas. A planilha mostra o nível de perigo e a ordem de pagamento (avalanche e bola de neve).",
    "4. 'Plano': quanto sobra por mês, qual dívida atacar primeiro e uma estimativa de prazo.",
    "5. 'Simulador de Acordo': compare até 2 propostas de acordo antes de aceitar.",
    "6. 'Rotativo': veja quanto uma dívida no rotativo cresce e o limite legal de 100% (Lei 14.690/2023).",
    "7. 'Gastos Anuais': IPVA, IPTU, material escolar... a planilha diz quanto guardar por mês.",
    "8. 'Reserva': sua meta de reserva de emergência e em quanto tempo chega lá.",
    "",
    "Preencha só as células AMARELAS. Datas no formato dia/mês/ano. Funciona no Excel e no Google Planilhas.",
    "Os valores já preenchidos são EXEMPLOS: troque pelos seus.",
    "SEGURANÇA: não coloque aqui nem na IA seu CPF, senhas ou números de cartão e conta.",
    "Ferramenta educativa. Estimativas não substituem os valores oficiais informados pelos credores.",
])

# ---------------- Orçamento ----------------
ob = wb.create_sheet("Orçamento")
t.titulo(ob, "Orçamento do mês")
ob["B2"] = "Mês de referência (1º dia do mês):"
ob["B2"].font = Font(bold=True)
ob["C2"] = "=DATE(YEAR(TODAY()),MONTH(TODAY()),1)"
t.entrada(ob["C2"], 'MM/YYYY')
t.cabecalho(ob, 4, ["Tipo", "Categoria", "Planejado (R$)", "Real (do Controle)", "Diferença", "Situação"])
linhas = ([("Entrada", c) for c in ["Salário / renda principal", "Renda extra", "Outras entradas"]]
          + [("Essencial", c) for c in ["Moradia (aluguel, prestação, condomínio)", "Contas da casa (luz, água, gás)",
                                         "Internet e celular", "Mercado", "Transporte", "Saúde e farmácia", "Educação"]]
          + [("Pessoal", c) for c in ["Alimentação fora e delivery", "Assinaturas", "Lazer", "Compras e roupas",
                                       "Cuidados pessoais", "Outros"]]
          + [("Futuro", c) for c in ["Parcelas de dívidas e acordos", "Reserva para gastos anuais", "Reserva de emergência"]])
exemplo = {"Salário / renda principal": 3800, "Renda extra": 600, "Moradia (aluguel, prestação, condomínio)": 1100, "Contas da casa (luz, água, gás)": 280,
           "Internet e celular": 150, "Mercado": 800, "Transporte": 250, "Saúde e farmácia": 100, "Alimentação fora e delivery": 150,
           "Assinaturas": 40, "Lazer": 80, "Reserva de emergência": 50}
F = 5
L = F + len(linhas) - 1
for i, (tipo, cat) in enumerate(linhas):
    r = F + i
    ob.cell(row=r, column=1, value=tipo)
    ob.cell(row=r, column=2, value=cat)
    c = ob.cell(row=r, column=3)
    if cat == "Parcelas de dívidas e acordos":
        c.value = "='Minhas Dívidas'!F18"
        t.calculo(c, BRL)
    elif cat == "Reserva para gastos anuais":
        c.value = "='Gastos Anuais'!D22"
        t.calculo(c, BRL)
    else:
        c.value = exemplo.get(cat)
        t.entrada(c, BRL)
    if tipo != "Entrada":
        ob.cell(row=r, column=4, value=(f"=SUMIFS('Controle de Gastos'!$D$5:$D$1004,'Controle de Gastos'!$C$5:$C$1004,B{r},"
                                        f"'Controle de Gastos'!$A$5:$A$1004,\">=\"&$C$2,'Controle de Gastos'!$A$5:$A$1004,\"<=\"&EOMONTH($C$2,0))"))
        ob.cell(row=r, column=5, value=f"=C{r}-D{r}")
        ob.cell(row=r, column=6, value=f'=IF(C{r}=0,IF(D{r}>0,"Não planejado",""),IF(D{r}>C{r},"Estourou",IF(D{r}>C{r}*0.8,"Atenção","Dentro")))')
        t.calculo(ob.cell(row=r, column=4), BRL)
        t.calculo(ob.cell(row=r, column=5), BRL)
        t.calculo(ob.cell(row=r, column=6), negrito=True)
vermelho_se(ob, f"F{F}:F{L}", f'OR($F{F}="Estourou",$F{F}="Não planejado")')
amarelo_se(ob, f"F{F}:F{L}", f'$F{F}="Atenção"')
verde_se(ob, f"F{F}:F{L}", f'$F{F}="Dentro"')
S = L + 2
res = [("Total de entradas", f'=SUMIF(A{F}:A{L},"Entrada",C{F}:C{L})', BRL),
       ("Total essencial", f'=SUMIF(A{F}:A{L},"Essencial",C{F}:C{L})', BRL),
       ("Total pessoal", f'=SUMIF(A{F}:A{L},"Pessoal",C{F}:C{L})', BRL),
       ("Total futuro (dívidas e reservas)", f'=SUMIF(A{F}:A{L},"Futuro",C{F}:C{L})', BRL),
       ("SOBRA (ou FALTA) DO MÊS", f"=C{S}-C{S+1}-C{S+2}-C{S+3}", BRL),
       ("% da renda em essenciais", f"=IF(C{S}=0,0,C{S+1}/C{S})", PCT),
       ("% da renda em dívidas", f"=IF(C{S}=0,0,'Minhas Dívidas'!F18/C{S})", PCT),
       ("Gasto real no mês (Controle)", f"=SUM(D{F}:D{L})", BRL)]
for i, (n, f, fmt) in enumerate(res):
    r = S + i
    ob.cell(row=r, column=2, value=n).font = Font(bold=True)
    t.calculo(ob.cell(row=r, column=3, value=f), fmt, True)
SOBRA = f"Orçamento!$C${S+4}"
ESSENCIAL = f"Orçamento!$C${S+1}"
vermelho_se(ob, f"C{S+4}", f"$C${S+4}<0")
t.nota(ob, f"B{S+9}", "Referências: essenciais até 50–60% da renda; dívidas acima de 30% da renda é sinal de alerta.")
larguras(ob, [11, 42, 16, 18, 14, 15])
CATS = f"Orçamento!$B${F+3}:$B${L}"

# ---------------- Controle de Gastos ----------------
cg = wb.create_sheet("Controle de Gastos")
t.titulo(cg, "Controle de gastos", "Registre cada gasto. Leva 2 minutos por dia e é o hábito que mais muda o seu mês.")
t.cabecalho(cg, 4, ["Data", "Descrição", "Categoria", "Valor (R$)", "Forma de pagamento"])
dvc = DataValidation(type="list", formula1=f"={CATS}", allow_blank=True)
dvp = DataValidation(type="list", formula1='"Pix,Débito,Crédito,Dinheiro,Boleto"', allow_blank=True)
cg.add_data_validation(dvc); cg.add_data_validation(dvp)
dvc.add("C5:C1004"); dvp.add("E5:E1004")
for r in range(5, 1005):
    t.entrada(cg.cell(row=r, column=1), DATA)
    t.entrada(cg.cell(row=r, column=2))
    t.entrada(cg.cell(row=r, column=3))
    t.entrada(cg.cell(row=r, column=4), BRL)
    t.entrada(cg.cell(row=r, column=5))
for i, (d, desc, cat, v, fp) in enumerate([("=TODAY()", "Mercado (exemplo)", "Mercado", 185.4, "Débito"),
                                            ("=TODAY()", "Delivery (exemplo)", "Alimentação fora e delivery", 42.9, "Crédito")]):
    for c, val in enumerate([d, desc, cat, v, fp], 1):
        cg.cell(row=5 + i, column=c, value=val)
cg["G4"], cg["G5"], cg["G6"] = "Total no mês de referência", "No crédito", "Nº de gastos"
t.calculo(cg.cell(row=4, column=8, value='=SUMIFS(D5:D1004,A5:A1004,">="&Orçamento!$C$2,A5:A1004,"<="&EOMONTH(Orçamento!$C$2,0))'), BRL, True)
t.calculo(cg.cell(row=5, column=8, value='=SUMIFS(D5:D1004,E5:E1004,"Crédito",A5:A1004,">="&Orçamento!$C$2,A5:A1004,"<="&EOMONTH(Orçamento!$C$2,0))'), BRL, True)
t.calculo(cg.cell(row=6, column=8, value='=COUNTIFS(A5:A1004,">="&Orçamento!$C$2,A5:A1004,"<="&EOMONTH(Orçamento!$C$2,0))'), '0', True)
larguras(cg, [12, 34, 38, 14, 18, 3, 26, 16])
cg.freeze_panes = "A5"

# ---------------- Minhas Dívidas ----------------
dv = wb.create_sheet("Minhas Dívidas")
t.titulo(dv, "Minhas dívidas", "Liste todas. Juros ao mês em %: ex. 12% ao mês = digite 12.")
cols = ["Credor", "Tipo", "Saldo devedor hoje", "Juros ao mês (%)", "Atrasada?", "Parcela / mínimo",
        "Ordem Avalanche\n(juros maiores 1º)", "Ordem Bola de Neve\n(menor saldo 1º)", "Nível de perigo", "Tem garantia?"]
t.cabecalho(dv, 4, cols, 40)
tipos = ('"Cartão - rotativo,Cheque especial,Cartão - parcelamento,Empréstimo pessoal,Crediário / loja,'
         'Consignado,Financiamento de veículo,Financiamento imobiliário,Conta de consumo atrasada,Outro"')
dvt = DataValidation(type="list", formula1=tipos, allow_blank=True)
dvs = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
dv.add_data_validation(dvt); dv.add_data_validation(dvs)
first, last = 5, 16
for r in range(first, last + 1):
    for c, fmt in [(1, None), (2, None), (3, BRL), (4, '0.00'), (5, None), (6, BRL), (10, None)]:
        t.entrada(dv.cell(row=r, column=c), fmt)
    dv.cell(row=r, column=7, value=(f'=IF(C{r}>0,COUNTIFS($C${first}:$C${last},">0",$D${first}:$D${last},">"&D{r})'
                                    f'+COUNTIFS($C${first}:C{r},">0",$D${first}:D{r},D{r}),"")'))
    dv.cell(row=r, column=8, value=(f'=IF(C{r}>0,COUNTIFS($C${first}:$C${last},">0",$C${first}:$C${last},"<"&C{r})'
                                    f'+COUNTIFS($C${first}:C{r},C{r}),"")'))
    dv.cell(row=r, column=9, value=(f'=IF(B{r}="","",IF(OR(B{r}="Cartão - rotativo",B{r}="Cheque especial"),"Muito alto",'
                                    f'IF(OR(B{r}="Cartão - parcelamento",B{r}="Empréstimo pessoal",B{r}="Crediário / loja"),"Alto",'
                                    f'IF(B{r}="Conta de consumo atrasada","Prioridade (serviço essencial)",'
                                    f'IF(OR(B{r}="Consignado",B{r}="Financiamento de veículo"),"Médio","Baixo em juros")))))'))
    for c in (7, 8, 9):
        cell = dv.cell(row=r, column=c)
        t.calculo(cell, negrito=(c == 9))
        cell.alignment = Alignment(horizontal="center")
dvt.add(f"B{first}:B{last}"); dvs.add(f"E{first}:E{last}"); dvs.add(f"J{first}:J{last}")
vermelho_se(dv, f"I{first}:I{last}", f'OR($I{first}="Muito alto",$I{first}="Prioridade (serviço essencial)")')
amarelo_se(dv, f"I{first}:I{last}", f'$I{first}="Alto"')
ex = [("Banco X (exemplo)", "Cartão - rotativo", 2300, 13, "Não", 350, "Não"),
      ("Loja Y (exemplo)", "Crediário / loja", 600, 6, "Sim", 120, "Não"),
      ("Banco Z (exemplo)", "Empréstimo pessoal", 4000, 4.5, "Não", 290, "Não")]
for i, row in enumerate(ex):
    for c, v in zip((1, 2, 3, 4, 5, 6, 10), row):
        dv.cell(row=first + i, column=c, value=v)
tr = 18
dv.cell(row=tr, column=2, value="TOTAIS").font = Font(bold=True)
t.calculo(dv.cell(row=tr, column=3, value=f"=SUM(C{first}:C{last})"), BRL, True)
t.calculo(dv.cell(row=tr, column=6, value=f"=SUM(F{first}:F{last})"), BRL, True)
dv.cell(row=tr + 1, column=2, value="Juros por mês (aprox.)").font = Font(bold=True)
t.calculo(dv.cell(row=tr + 1, column=3, value=f"=SUMPRODUCT(C{first}:C{last},D{first}:D{last})/100"), BRL, True)
t.nota(dv, f"A{tr+3}", "'Juros por mês' é quanto as suas dívidas crescem por mês se nada for pago. Apague as linhas de exemplo ao começar.")
larguras(dv, [24, 26, 16, 12, 11, 15, 16, 16, 26, 12])
dv.freeze_panes = "A5"

# ---------------- Plano ----------------
pl = wb.create_sheet("Plano")
t.titulo(pl, "Plano de quitação", "Resumo automático. Use o Prompt 14 para o calendário mês a mês.")
D = "'Minhas Dívidas'"
R = f"{D}!$A${first}:$A${last}"
items = [("Sobra do mês (Orçamento)", f"={SOBRA}", BRL, False),
         ("Quanto vou colocar A MAIS nas dívidas por mês", 100, BRL, True),
         ("Total que devo hoje", f"={D}!C{tr}", BRL, False),
         ("Soma das parcelas e mínimos", f"={D}!F{tr}", BRL, False),
         ("Juros por mês (aprox.)", f"={D}!C{tr+1}", BRL, False),
         ("AVALANCHE: atacar primeiro", f'=IFERROR(INDEX({R},MATCH(1,{D}!$G${first}:$G${last},0)),"—")', None, False),
         ("BOLA DE NEVE: atacar primeiro", f'=IFERROR(INDEX({R},MATCH(1,{D}!$H${first}:$H${last},0)),"—")', None, False),
         ("Pagamento mensal total nas dívidas", "=B7+B5", BRL, False),
         ("O pagamento cobre os juros?", '=IF(B11<=B8,"NÃO: a dívida cresce. Negocie ou corte gastos","Sim")', None, False),
         ("Prazo estimado para quitar tudo (meses, aprox.)",
          '=IF(B11<=B8,"—",IFERROR(ROUNDUP(-LN(1-B6*(B8/B6)/B11)/LN(1+B8/B6),0),"—"))', '0', False)]
for i, (n, f, fmt, entrada) in enumerate(items):
    r = 4 + i
    pl.cell(row=r, column=1, value=n).font = Font(bold=True)
    c = pl.cell(row=r, column=2, value=f)
    (t.entrada if entrada else t.calculo)(c, fmt)
vermelho_se(pl, "B12", 'LEFT($B$12,3)="NÃO"')
vermelho_se(pl, "B4", "$B$4<0")
t.nota(pl, "A15", "O prazo usa a taxa média de juros de todas as dívidas e é só uma referência. O plano real (com a ordem de pagamento) vem do Prompt 14.")
t.nota(pl, "A16", "Se a sobra é negativa: primeiro corte gastos (Prompt 3) e negocie parcelas menores (Prompt 18). Se nem assim fecha, leia o capítulo de superendividamento.")
larguras(pl, [52, 44])

# ---------------- Simulador de Acordo ----------------
sa = wb.create_sheet("Simulador de Acordo")
t.titulo(sa, "Simulador de acordo", "Compare propostas antes de aceitar. Só aceite parcela que caiba com folga.")
t.cabecalho(sa, 4, ["", "Proposta A", "Proposta B"])
ent = [("Dívida original (valor que você devia)", 3000, 3000, BRL), ("Entrada (R$)", 150, 0, BRL),
       ("Número de parcelas", 10, 1, '0'), ("Valor de cada parcela (R$)", 150, 1200, BRL)]
for i, (n, a, b, fmt) in enumerate(ent):
    r = 5 + i
    sa.cell(row=r, column=1, value=n)
    t.entrada(sa.cell(row=r, column=2, value=a), fmt)
    t.entrada(sa.cell(row=r, column=3, value=b), fmt)
calc = [("TOTAL QUE VOU PAGAR", "=B6+B7*B8", BRL), ("Desconto (ou custo extra) em R$", "=B5-B9", BRL),
        ("Desconto em %", "=IF(B5=0,0,B10/B5)", PCT), ("Maior pagamento em um mês", "=MAX(B6,B8)", BRL),
        ("Cabe no orçamento?", f'=IF(B12<={SOBRA}*0.7,"Cabe com folga",IF(B12<={SOBRA},"Cabe, mas apertado","NÃO CABE"))', None)]
for i, (n, f, fmt) in enumerate(calc):
    r = 9 + i
    sa.cell(row=r, column=1, value=n).font = Font(bold=True)
    t.calculo(sa.cell(row=r, column=2, value=f), fmt, True)
    t.calculo(sa.cell(row=r, column=3, value=f.replace("B", "C").replace("$C$", "$B$").replace("Orçamento!$C$", "Orçamento!$C$")), fmt, True)
sa["C13"] = f'=IF(C12<={SOBRA}*0.7,"Cabe com folga",IF(C12<={SOBRA},"Cabe, mas apertado","NÃO CABE"))'
sa["A15"] = "Sobra do mês considerada (aba Orçamento)"
t.calculo(sa.cell(row=15, column=2, value=f"={SOBRA}"), BRL)
vermelho_se(sa, "B13:C13", 'B13="NÃO CABE"')
amarelo_se(sa, "B13:C13", 'B13="Cabe, mas apertado"')
verde_se(sa, "B13:C13", 'B13="Cabe com folga"')
t.nota(sa, "A17", "Pergunte sempre: valor total, CET, multa se atrasar uma parcela e quando o nome sai do cadastro negativo. Peça tudo por escrito.")
larguras(sa, [42, 18, 18])

# ---------------- Rotativo ----------------
ro = wb.create_sheet("Rotativo")
t.titulo(ro, "Quanto o rotativo cresce", "Simulação com juros compostos e o teto legal de 100% da dívida original (Lei 14.690/2023).")
for i, (n, v, fmt) in enumerate([("Valor que entrou no rotativo (R$)", 1000, BRL), ("Juros ao mês (%)", 0.14, PCT),
                                  ("Quanto você paga por mês (R$)", 0, BRL)]):
    ro.cell(row=4 + i, column=1, value=n)
    t.entrada(ro.cell(row=4 + i, column=2, value=v), fmt)
ro["A7"] = "Teto legal da dívida (original + 100%)"
ro["A7"].font = Font(bold=True)
t.calculo(ro.cell(row=7, column=2, value="=B4*2"), BRL, True)
t.cabecalho(ro, 9, ["Mês", "Dívida sem o teto", "Dívida com o teto legal", "Juros acumulados (com teto)"])
for m in range(0, 13):
    r = 10 + m
    ro.cell(row=r, column=1, value=m)
    if m == 0:
        ro.cell(row=r, column=2, value="=B4")
        ro.cell(row=r, column=3, value="=B4")
    else:
        ro.cell(row=r, column=2, value=f"=MAX(0,B{r-1}*(1+$B$5)-$B$6)")
        ro.cell(row=r, column=3, value=f"=MAX(0,MIN(C{r-1}*(1+$B$5),$B$7)-$B$6)")
    ro.cell(row=r, column=4, value=f"=MAX(0,C{r}-$B$4+$B$6*A{r})")
    for c in (2, 3, 4):
        t.calculo(ro.cell(row=r, column=c), BRL)
t.nota(ro, "A24", "Simplificação para fins educativos: o teto considera juros e encargos sobre o valor original. Na prática, o banco deve oferecer parcelamento após 1 fatura.")
t.nota(ro, "A25", "Se o banco cobra mais do que o teto, peça o demonstrativo e reclame (Prompt 12 e Prompt 22).")
larguras(ro, [40, 20, 24, 26])

# ---------------- Gastos Anuais ----------------
ga = wb.create_sheet("Gastos Anuais")
t.titulo(ga, "Gastos anuais (os escondidos)", "Quanto guardar por mês para eles não virarem dívida.")
t.cabecalho(ga, 4, ["Gasto", "Mês em que vence", "Valor no ano (R$)", "Guardar por mês"])
itens = [("IPVA e licenciamento", "Jan–Mar", 1200), ("IPTU", "Fev–Mar", 900), ("Material e uniforme escolar", "Jan–Fev", 600),
         ("Matrícula escolar", "Dez–Jan", 0), ("Seguro do carro", "", 0), ("Manutenção do carro", "", 600),
         ("Presentes de fim de ano", "Dez", 500), ("Aniversários", "Vários", 300), ("Viagem / férias", "", 0),
         ("Anuidades", "", 0), ("Vacinas e veterinário (pets)", "", 0), ("Consertos na casa", "", 400),
         ("Óculos, dentista e exames", "", 300), ("Outros", "", 0)]
for i in range(16):
    r = 5 + i
    for c, fmt in [(1, None), (2, None), (3, BRL)]:
        t.entrada(ga.cell(row=r, column=c), fmt)
    if i < len(itens):
        for c, v in enumerate(itens[i], 1):
            ga.cell(row=r, column=c, value=v)
    ga.cell(row=r, column=4, value=f'=IF(C{r}="","",C{r}/12)')
    t.calculo(ga.cell(row=r, column=4), BRL)
ga["A22"], ga["C22"] = "TOTAL", "=SUM(C5:C20)"
ga["A22"].font = Font(bold=True)
t.calculo(ga["C22"], BRL, True)
t.calculo(ga.cell(row=22, column=4, value="=C22/12"), BRL, True)
t.nota(ga, "A24", "O total por mês vai automaticamente para a linha 'Reserva para gastos anuais' do Orçamento.")
larguras(ga, [34, 16, 18, 16])

# ---------------- Reserva ----------------
rs = wb.create_sheet("Reserva")
t.titulo(rs, "Reserva de emergência", "Comece pequeno. O objetivo final é de 3 a 6 meses de gastos essenciais.")
ent = [("Gastos essenciais por mês (do Orçamento)", f"={ESSENCIAL}", BRL, False), ("Meta em meses de gastos", 3, '0', True),
       ("Quanto já tenho guardado", 0, BRL, True), ("Quanto consigo guardar por mês", 100, BRL, True)]
for i, (n, v, fmt, e) in enumerate(ent):
    rs.cell(row=4 + i, column=1, value=n)
    (t.entrada if e else t.calculo)(rs.cell(row=4 + i, column=2, value=v), fmt)
res = [("1ª meta: R$ 500 ou 1 mês de essenciais (o menor)", "=MIN(500,B4)", BRL),
       ("Meses para a 1ª meta", '=IF(B7<=0,"—",MAX(0,ROUNDUP((B9-B6)/B7,0)))', '0'),
       ("META FINAL", "=B4*B5", BRL), ("Falta para a meta final", "=MAX(0,B11-B6)", BRL),
       ("Meses para a meta final", '=IF(B7<=0,"—",ROUNDUP(B12/B7,0))', '0'), ("Progresso", "=IF(B11=0,0,MIN(1,B6/B11))", PCT)]
for i, (n, f, fmt) in enumerate(res):
    rs.cell(row=9 + i, column=1, value=n).font = Font(bold=True)
    t.calculo(rs.cell(row=9 + i, column=2, value=f), fmt, True)
t.nota(rs, "A16", "Guarde em aplicação segura e com resgate imediato. Emergência é o que é urgente E necessário: saúde, perda de renda, conserto essencial.")
larguras(rs, [52, 18])

wb.save(sys.argv[1])
print("ok", sys.argv[1])
