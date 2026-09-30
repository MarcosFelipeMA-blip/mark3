import sys
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment
from estilo import Tema, BRL, PCT, DATA, larguras, vermelho_se, verde_se, amarelo_se, como_usar

t = Tema("1F5F8B", "EEF4F9")
wb = Workbook()

como_usar(wb.active, t, "Planilhas do Corretor", [
    "COMO USAR",
    "1. 'Meta Reversa': diga quanto quer ganhar e suas taxas de conversão. Descubra quantos leads, atendimentos e visitas precisa por semana.",
    "2. 'Funil': cadastre cada lead. Na coluna Etapa use 1=Lead, 2=Atendimento, 3=Visita, 4=Proposta, 5=Venda. Atualize o 'Último contato'.",
    "   A coluna 'Próximo contato' segue a cadência de acompanhamento. Vermelho = atrasado. Amarelo = hoje.",
    "3. 'Painel': mostra suas taxas reais de conversão e onde está o gargalo. Depois de 2 ou 3 meses, use essas taxas na Meta Reversa.",
    "4. 'ACM': análise comparativa de mercado. Coloque imóveis parecidos e veja o preço por m² e a faixa de preço sugerida.",
    "5. 'Captações': controle prazos de exclusividade e imóveis que precisam de revisão de preço.",
    "6. 'Comissões': o que você já recebeu e o que ainda vai receber.",
    "7. 'Simuladores': custos da compra, SAC x Price e rentabilidade de aluguel. São ESTIMATIVAS: os valores oficiais vêm do banco, prefeitura e cartório.",
    "8. 'Calendário 30 dias': o que postar em cada dia, com o prompt da Biblioteca.",
    "",
    "Preencha só as células AMARELAS. Datas no formato dia/mês/ano. Funciona no Excel e no Google Planilhas.",
    "Os valores já preenchidos são EXEMPLOS: troque pelos seus.",
])

# ---------------- Meta Reversa ----------------
mr = wb.create_sheet("Meta Reversa")
t.titulo(mr, "Meta reversa", "Comece pelo quanto quer ganhar e descubra o que precisa fazer por semana.")
ent = [("Quanto quero ganhar por mês (R$)", 10000, BRL), ("Ticket médio dos imóveis (R$)", 500000, BRL),
       ("Comissão total da venda (%)", 0.06, PCT), ("Minha parte da comissão (%)", 0.50, PCT),
       ("Lead → atendimento (%)", 0.50, PCT), ("Atendimento → visita (%)", 0.33, PCT),
       ("Visita → proposta (%)", 0.20, PCT), ("Proposta → venda (%)", 0.25, PCT)]
for i, (n, v, fmt) in enumerate(ent):
    mr.cell(row=4 + i, column=1, value=n)
    t.entrada(mr.cell(row=4 + i, column=2, value=v), fmt)
t.nota(mr, "C8", "← taxas de referência para começar. Use as suas reais do 'Painel'.")
t.cabecalho(mr, 13, ["Resultado", "Por mês", "Por semana"])
res = [("Minha comissão por venda", "=B5*B6*B7", None, BRL),
       ("Vendas necessárias", "=IF(B14=0,0,B4/B14)", "=B15/4.3", '0.00'),
       ("Propostas necessárias", "=IF(B11=0,0,B15/B11)", "=B16/4.3", '0.0'),
       ("Visitas necessárias", "=IF(B10=0,0,B16/B10)", "=B17/4.3", '0.0'),
       ("Atendimentos necessários", "=IF(B9=0,0,B17/B9)", "=B18/4.3", '0.0'),
       ("Leads necessários", "=IF(B8=0,0,B18/B8)", "=B19/4.3", '0')]
for i, (n, f1, f2, fmt) in enumerate(res):
    r = 14 + i
    mr.cell(row=r, column=1, value=n).font = Font(bold=True)
    t.calculo(mr.cell(row=r, column=2, value=f1), fmt, True)
    if f2:
        t.calculo(mr.cell(row=r, column=3, value=f2), fmt, True)
t.nota(mr, "A21", "Leitura: 0,67 venda por mês = 2 vendas a cada 3 meses. Se os leads necessários forem muitos, melhore a taxa do gargalo antes de buscar mais leads.")
larguras(mr, [38, 18, 16])

# ---------------- Funil ----------------
fu = wb.create_sheet("Funil")
t.titulo(fu, "Funil de leads", "Etapa: 1=Lead  2=Atendimento  3=Visita  4=Proposta  5=Venda.  Marque 'Perdido' = Sim quando o lead desistir.")
fu["M1"], fu["M2"] = "Cadência (dias até o próximo contato)", "Etapa → dias"
cad = [(1, 1), (2, 3), (3, 2), (4, 1), (5, 30)]
for i, (e, d) in enumerate(cad):
    fu.cell(row=3 + i, column=13, value=e)
    t.entrada(fu.cell(row=3 + i, column=14, value=d), '0')
cols = ["Entrada", "Nome", "WhatsApp", "Origem", "Finalidade", "Interesse (tipo, bairro, faixa)",
        "Etapa (1-5)", "Etapa", "Último contato", "Próximo contato", "Situação", "Perdido?", "Motivo da perda / notas"]
t.cabecalho(fu, 9, cols, 36)
F, L = 10, 509
dvo = DataValidation(type="list", formula1='"Portal,Instagram,Indicação,Placa,Site,WhatsApp,Parceiro,Outro"', allow_blank=True)
dvf = DataValidation(type="list", formula1='"Morar,Investir,Alugar,Vender"', allow_blank=True)
dve = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True)
dvp = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
for d in (dvo, dvf, dve, dvp):
    fu.add_data_validation(d)
dvo.add(f"D{F}:D{L}"); dvf.add(f"E{F}:E{L}"); dve.add(f"G{F}:G{L}"); dvp.add(f"L{F}:L{L}")
for r in range(F, L + 1):
    for c in (2, 3, 4, 5, 6, 7, 12, 13):
        t.entrada(fu.cell(row=r, column=c))
    t.entrada(fu.cell(row=r, column=1), DATA)
    t.entrada(fu.cell(row=r, column=9), DATA)
    fu.cell(row=r, column=8, value=f'=IF(G{r}="","",CHOOSE(G{r},"Lead","Atendimento","Visita","Proposta","Venda"))')
    fu.cell(row=r, column=10, value=(f'=IF(OR(G{r}="",I{r}="",L{r}="Sim"),"",I{r}+VLOOKUP(G{r},$M$3:$N$7,2,FALSE))'))
    fu.cell(row=r, column=11, value=(f'=IF(L{r}="Sim","Perdido",IF(J{r}="","",IF(J{r}<TODAY(),"ATRASADO",'
                                     f'IF(J{r}=TODAY(),"Hoje","Em dia"))))'))
    t.calculo(fu.cell(row=r, column=8))
    t.calculo(fu.cell(row=r, column=10), DATA)
    t.calculo(fu.cell(row=r, column=11), negrito=True)
vermelho_se(fu, f"K{F}:K{L}", f'$K{F}="ATRASADO"')
amarelo_se(fu, f"K{F}:K{L}", f'$K{F}="Hoje"')
verde_se(fu, f"K{F}:K{L}", f'$K{F}="Em dia"')
exemplos = [(-12, "Rafael (exemplo)", "Portal", "Morar", "2 dorms Tatuapé até 550 mil", 3, -5, "Não"),
            (-3, "Juliana (exemplo)", "Instagram", "Morar", "3 dorms Mooca", 2, -1, "Não"),
            (-20, "Marcos (exemplo)", "Indicação", "Investir", "Studio perto do metrô", 4, -1, "Não")]
for i, (de, n, o, fi, it, e, uc, p) in enumerate(exemplos):
    r = F + i
    fu.cell(row=r, column=1, value=f"=TODAY(){de}")
    fu.cell(row=r, column=2, value=n); fu.cell(row=r, column=4, value=o); fu.cell(row=r, column=5, value=fi)
    fu.cell(row=r, column=6, value=it); fu.cell(row=r, column=7, value=e)
    fu.cell(row=r, column=9, value=f"=TODAY(){uc}"); fu.cell(row=r, column=12, value=p)
larguras(fu, [12, 22, 16, 12, 11, 30, 9, 13, 13, 13, 12, 9, 30, 12, 8])
fu.freeze_panes = "C10"
fu.auto_filter.ref = f"A9:M{L}"

# ---------------- Painel ----------------
pa = wb.create_sheet("Painel")
t.titulo(pa, "Painel do funil", "Calculado a partir da aba Funil. Mostra suas taxas reais e o gargalo.")
t.cabecalho(pa, 4, ["Etapa", "Chegaram a esta etapa", "Conversão da etapa anterior"])
R = f"Funil!$G${F}:$G${L}"
etapas = [("Leads", f'=COUNTIF({R},">=1")', None), ("Atendimentos", f'=COUNTIF({R},">=2")', "=IF(B5=0,\"\",B6/B5)"),
          ("Visitas", f'=COUNTIF({R},">=3")', "=IF(B6=0,\"\",B7/B6)"),
          ("Propostas", f'=COUNTIF({R},">=4")', "=IF(B7=0,\"\",B8/B7)"),
          ("Vendas", f'=COUNTIF({R},5)', "=IF(B8=0,\"\",B9/B8)")]
for i, (n, f1, f2) in enumerate(etapas):
    r = 5 + i
    pa.cell(row=r, column=1, value=n).font = Font(bold=True)
    t.calculo(pa.cell(row=r, column=2, value=f1), '0', True)
    if f2:
        t.calculo(pa.cell(row=r, column=3, value=f2), PCT, True)
pa["A11"], pa["A12"], pa["A13"] = "Leads atrasados (sem contato)", "Contatos para hoje", "Leads perdidos"
t.calculo(pa.cell(row=11, column=2, value=f'=COUNTIF(Funil!$K${F}:$K${L},"ATRASADO")'), '0', True)
t.calculo(pa.cell(row=12, column=2, value=f'=COUNTIF(Funil!$K${F}:$K${L},"Hoje")'), '0', True)
t.calculo(pa.cell(row=13, column=2, value=f'=COUNTIF(Funil!$L${F}:$L${L},"Sim")'), '0', True)
pa["A15"] = "GARGALO (menor conversão)"
pa["A15"].font = Font(bold=True)
t.calculo(pa.cell(row=15, column=2, value=(
    '=IF(COUNT(C6:C9)<4,"Registre mais leads para calcular",INDEX(A6:A9,MATCH(MIN(C6:C9),C6:C9,0)))')), negrito=True)
t.nota(pa, "A17", "Gargalo em Atendimentos: responda mais rápido. Em Visitas: qualifique melhor. Em Propostas: melhore a visita. Em Vendas: preço e negociação.")
larguras(pa, [34, 26, 28])

# ---------------- ACM ----------------
ac = wb.create_sheet("ACM")
t.titulo(ac, "Análise comparativa de mercado (ACM)", "Use de 5 a 10 imóveis parecidos. Marque 'Não' em 'Usar?' para descartar extremos.")
ac["A4"], ac["A5"], ac["A6"] = "Imóvel avaliado", "Área (m²)", "Ajuste (%) para diferenças: andar, vaga, reforma, sol..."
t.entrada(ac["B4"]); ac["B4"] = "Apto 2 dorms, Rua X (exemplo)"
t.entrada(ac["B5"], '0'); ac["B5"] = 65
t.entrada(ac["B6"], PCT); ac["B6"] = 0
t.nota(ac, "C6", "Ex.: +5% se o avaliado é reformado e os demais não; −5% se não tem vaga.")
t.cabecalho(ac, 8, ["Endereço / condomínio", "Área (m²)", "Preço (R$)", "Fonte", "Usar?", "Preço por m²", "Considerado"])
dvfonte = DataValidation(type="list", formula1='"Anúncio,Venda real"', allow_blank=True)
dvusar = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ac.add_data_validation(dvfonte); ac.add_data_validation(dvusar)
comp = [("Cond. A", 62, 520000), ("Cond. B", 70, 560000), ("Cond. C", 64, 505000), ("Cond. D", 68, 590000), ("Cond. E", 60, 470000)]
for i in range(12):
    r = 9 + i
    for c, fmt in [(1, None), (2, '0'), (3, BRL), (4, None), (5, None)]:
        t.entrada(ac.cell(row=r, column=c), fmt)
    if i < len(comp):
        ac.cell(row=r, column=1, value=comp[i][0] + " (exemplo)")
        ac.cell(row=r, column=2, value=comp[i][1]); ac.cell(row=r, column=3, value=comp[i][2])
        ac.cell(row=r, column=4, value="Anúncio"); ac.cell(row=r, column=5, value="Sim")
    ac.cell(row=r, column=6, value=f'=IF(OR(B{r}="",C{r}="",B{r}=0),"",C{r}/B{r})')
    ac.cell(row=r, column=7, value=f'=IF(AND(E{r}="Sim",F{r}<>""),F{r},"")')
    t.calculo(ac.cell(row=r, column=6), BRL); t.calculo(ac.cell(row=r, column=7), BRL)
dvfonte.add("D9:D20"); dvusar.add("E9:E20")
out = [("Comparáveis considerados", "=COUNT(G9:G20)", '0'), ("Menor preço por m²", "=IF(B23=0,\"\",MIN(G9:G20))", BRL),
       ("Maior preço por m²", "=IF(B23=0,\"\",MAX(G9:G20))", BRL), ("Média do m²", "=IF(B23=0,\"\",AVERAGE(G9:G20))", BRL),
       ("Mediana do m² (mais confiável)", "=IF(B23=0,\"\",MEDIAN(G9:G20))", BRL),
       ("VALOR DE REFERÊNCIA", "=IF(B23=0,\"\",B27*B5*(1+B6))", BRL),
       ("Faixa sugerida: de", "=IF(B23=0,\"\",B28*0.95)", BRL), ("Faixa sugerida: até", "=IF(B23=0,\"\",B28*1.05)", BRL)]
for i, (n, f, fmt) in enumerate(out):
    r = 23 + i
    ac.cell(row=r, column=1, value=n).font = Font(bold=True)
    t.calculo(ac.cell(row=r, column=2, value=f), fmt, True)
t.nota(ac, "A32", "Preço de anúncio costuma ter margem de negociação. Sempre que tiver valores de venda real, prefira-os. Leve o resultado ao Prompt 7.")
larguras(ac, [40, 18, 16, 12, 8, 14, 14])

# ---------------- Captações ----------------
cp = wb.create_sheet("Captações")
t.titulo(cp, "Captações e exclusividades", "Alertas automáticos: exclusividade vencendo e imóvel parado que precisa de revisão de preço.")
t.cabecalho(cp, 4, ["Imóvel", "Proprietário", "Bairro", "Preço anunciado", "Data da captação", "Exclusivo?", "Prazo (dias)",
                    "Exclusividade até", "Dias anunciado", "Visitas", "Propostas", "Alerta"], 36)
dvx = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
cp.add_data_validation(dvx); dvx.add("F5:F104")
for r in range(5, 105):
    for c, fmt in [(1, None), (2, None), (3, None), (4, BRL), (5, DATA), (6, None), (7, '0'), (10, '0'), (11, '0')]:
        t.entrada(cp.cell(row=r, column=c), fmt)
    cp.cell(row=r, column=8, value=f'=IF(OR(F{r}<>"Sim",E{r}="",G{r}=""),"",E{r}+G{r})')
    cp.cell(row=r, column=9, value=f'=IF(E{r}="","",TODAY()-E{r})')
    cp.cell(row=r, column=12, value=(
        f'=IF(E{r}="","",IF(AND(H{r}<>"",H{r}<TODAY()),"Exclusividade vencida",IF(AND(H{r}<>"",H{r}-TODAY()<=15),"Renovar exclusividade",'
        f'IF(AND(I{r}>=30,N(J{r})<3),"Revisar preço (poucas visitas)",IF(AND(I{r}>=60,N(K{r})=0),"Revisar preço (sem propostas)","OK")))))'))
    t.calculo(cp.cell(row=r, column=8), DATA); t.calculo(cp.cell(row=r, column=9), '0'); t.calculo(cp.cell(row=r, column=12), negrito=True)
cp["A5"], cp["C5"], cp["D5"], cp["E5"], cp["F5"], cp["G5"], cp["J5"], cp["K5"] = "Apto 2 dorms (exemplo)", "Tatuapé", 540000, "=TODAY()-40", "Sim", 90, 2, 0
vermelho_se(cp, "L5:L104", 'AND($L5<>"",$L5<>"OK")')
verde_se(cp, "L5:L104", '$L5="OK"')
larguras(cp, [26, 20, 14, 15, 14, 10, 10, 15, 10, 9, 10, 30])

# ---------------- Comissões ----------------
co = wb.create_sheet("Comissões")
t.titulo(co, "Comissões", "Controle do que foi recebido e do que está a receber.")
t.cabecalho(co, 4, ["Data da venda", "Imóvel", "Valor da venda", "Comissão (%)", "Minha parte (%)",
                    "Comissão total", "Minha comissão", "Recebido?"], 36)
dvr = DataValidation(type="list", formula1='"Sim,Não,Parcial"', allow_blank=True)
co.add_data_validation(dvr); dvr.add("H5:H104")
for r in range(5, 105):
    for c, fmt in [(1, DATA), (2, None), (3, BRL), (4, PCT), (5, PCT), (8, None)]:
        t.entrada(co.cell(row=r, column=c), fmt)
    co.cell(row=r, column=6, value=f'=IF(C{r}="","",C{r}*D{r})')
    co.cell(row=r, column=7, value=f'=IF(C{r}="","",F{r}*E{r})')
    t.calculo(co.cell(row=r, column=6), BRL); t.calculo(co.cell(row=r, column=7), BRL, True)
co["J4"], co["J5"], co["J6"], co["J7"] = "Total de vendas (VGV)", "Minha comissão total", "Recebido", "A receber"
t.calculo(co.cell(row=4, column=11, value="=SUM(C5:C104)"), BRL, True)
t.calculo(co.cell(row=5, column=11, value="=SUM(G5:G104)"), BRL, True)
t.calculo(co.cell(row=6, column=11, value='=SUMIF(H5:H104,"Sim",G5:G104)'), BRL, True)
t.calculo(co.cell(row=7, column=11, value="=K5-K6"), BRL, True)
larguras(co, [13, 30, 16, 12, 13, 15, 16, 11, 3, 24, 18])

# ---------------- Simuladores ----------------
si = wb.create_sheet("Simuladores")
t.titulo(si, "Simuladores (estimativas)", "Valores aproximados para orientar a conversa. Os oficiais vêm do banco, da prefeitura e do cartório.")
t.rotulo(si, "A4", "1. CUSTOS DA COMPRA")
cc = [("Valor do imóvel", 500000, BRL), ("Alíquota do ITBI da cidade (%)", 0.03, PCT),
      ("Escritura + registro (estimativa, % do valor)", 0.015, PCT), ("Avaliação e taxas do banco (R$)", 3500, BRL)]
for i, (n, v, fmt) in enumerate(cc):
    si.cell(row=5 + i, column=1, value=n); t.entrada(si.cell(row=5 + i, column=2, value=v), fmt)
for i, (n, f, fmt) in enumerate([("ITBI estimado", "=B5*B6", BRL), ("Escritura e registro estimados", "=B5*B7", BRL),
                                 ("TOTAL DE CUSTOS", "=B9+B10+B8", BRL), ("% do valor do imóvel", "=IF(B5=0,0,B11/B5)", PCT)]):
    si.cell(row=9 + i, column=1, value=n).font = Font(bold=True)
    t.calculo(si.cell(row=9 + i, column=2, value=f), fmt, True)
t.nota(si, "C6", "Confirme a alíquota e a base de cálculo na prefeitura.")

t.rotulo(si, "A15", "2. FINANCIAMENTO: SAC x PRICE")
fi = [("Valor do imóvel", 500000, BRL), ("Entrada (%)", 0.20, PCT), ("Juros ao ano (%)", 0.115, PCT), ("Prazo (meses)", 360, '0')]
for i, (n, v, fmt) in enumerate(fi):
    si.cell(row=16 + i, column=1, value=n); t.entrada(si.cell(row=16 + i, column=2, value=v), fmt)
calc = [("Valor da entrada", "=B16*B17", BRL), ("Valor financiado", "=B16-B20", BRL),
        ("Juros ao mês (equivalente)", "=(1+B18)^(1/12)-1", '0.000%'),
        ("PRICE: parcela fixa", "=IF(B22=0,B21/B19,B21*B22/(1-(1+B22)^(-B19)))", BRL),
        ("SAC: primeira parcela", "=B21/B19+B21*B22", BRL), ("SAC: última parcela", "=B21/B19*(1+B22)", BRL),
        ("PRICE: total de juros", "=B23*B19-B21", BRL), ("SAC: total de juros", "=B22*B21*(B19+1)/2", BRL),
        ("Renda mínima estimada (parcela = 30% da renda) PRICE", "=B23/0.3", BRL),
        ("Renda mínima estimada (parcela = 30% da renda) SAC", "=B24/0.3", BRL)]
for i, (n, f, fmt) in enumerate(calc):
    si.cell(row=20 + i, column=1, value=n).font = Font(bold=True)
    t.calculo(si.cell(row=20 + i, column=2, value=f), fmt, True)
t.nota(si, "A31", "Sem seguros obrigatórios, taxas administrativas e correção monetária. O banco informa o CET oficial.")

t.rotulo(si, "A34", "3. RENTABILIDADE DE ALUGUEL (investidor)")
al = [("Preço do imóvel", 350000, BRL), ("Aluguel mensal", 2000, BRL), ("Custos mensais do proprietário (IPTU, cond. se vago etc.)", 150, BRL),
      ("Taxa de administração (%)", 0.08, PCT), ("Meses vagos por ano (estimativa)", 1, '0')]
for i, (n, v, fmt) in enumerate(al):
    si.cell(row=35 + i, column=1, value=n); t.entrada(si.cell(row=35 + i, column=2, value=v), fmt)
for i, (n, f, fmt) in enumerate([("Rentabilidade bruta ao ano", "=IF(B35=0,0,B36*12/B35)", PCT),
                                 ("Renda líquida no ano", "=B36*(12-B39)*(1-B38)-B37*12", BRL),
                                 ("Rentabilidade líquida ao ano", "=IF(B35=0,0,B41/B35)", PCT),
                                 ("Rentabilidade líquida ao mês", "=(1+B42)^(1/12)-1", '0.00%')]):
    si.cell(row=40 + i, column=1, value=n).font = Font(bold=True)
    t.calculo(si.cell(row=40 + i, column=2, value=f), fmt, True)
t.nota(si, "A45", "Antes de impostos sobre aluguel e sem considerar valorização. Não é recomendação de investimento.")
larguras(si, [56, 18, 50])

# ---------------- Calendário ----------------
ca = wb.create_sheet("Calendário 30 dias")
t.titulo(ca, "Calendário de conteúdo: 30 dias", "Gere a semana inteira no domingo com o Prompt 40.")
t.cabecalho(ca, 4, ["Dia", "Semana / foco", "Dia da semana", "Tipo", "Ideia", "Prompt", "Feito?"])
plano = {"Seg": ("Imóvel", "Tour em vídeo do imóvel da semana", "6"), "Ter": ("Educação", "Dica para comprador (financiamento, custos, documentos)", "30"),
         "Qua": ("Bairro", "Guia do bairro em carrossel", "41"), "Qui": ("Imóvel", "Imóvel no feed com legenda", "5"),
         "Sex": ("Mercado", "Preço por m² ou conteúdo para proprietários", "43"), "Sáb": ("Prova", "Vendido, chave entregue ou depoimento", "42"),
         "Dom": ("Planejamento", "Gerar os conteúdos da semana", "40")}
focos = ["1 · Apresentação e território", "2 · Compradores", "3 · Proprietários (captação)", "4 · Prova e indicação", "5 · Revisão"]
dias = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
dvs = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ca.add_data_validation(dvs)
for d in range(30):
    r = 5 + d
    ds = dias[d % 7]
    for c, v in enumerate([d + 1, focos[d // 7], ds] + list(plano[ds]), 1):
        ca.cell(row=r, column=c, value=v)
    t.entrada(ca.cell(row=r, column=7)); dvs.add(f"G{r}")
verde_se(ca, "A5:G34", '$G5="Sim"')
larguras(ca, [6, 30, 14, 13, 52, 9, 9])

wb.save(sys.argv[1])
print("ok", sys.argv[1])
