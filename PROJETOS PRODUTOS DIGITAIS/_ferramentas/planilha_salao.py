import sys
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, Alignment
from estilo import Tema, BRL, PCT, DATA, larguras, vermelho_se, verde_se, amarelo_se, como_usar

t = Tema("B0476B", "FBF1F4")
wb = Workbook()

# ---------------- Como usar ----------------
como_usar(wb.active, t, "Planilhas do Salão", [
    "COMO USAR",
    "1. Aba 'Metas': preencha a estrutura do salão e descubra sua ocupação e quantos atendimentos por dia você precisa.",
    "2. Aba 'Precificação': coloque seus custos fixos e seus serviços. Os serviços em VERMELHO estão abaixo do preço mínimo.",
    "3. Aba 'Ciclos': confira de quantos em quantos dias cada serviço deve voltar (já vem preenchida, ajuste se quiser).",
    "4. Aba 'Clientes': cadastre suas clientes com a data da última visita. A coluna Status diz quem chamar hoje.",
    "   • Em dia: ainda não é hora.  • Chamar agora: passou do prazo (use o Prompt 28).  • Sumida: 15 a 60 dias de atraso (Prompts 29 e 30).  • Reativar: mais de 60 dias (Prompt 31).",
    "5. Aba 'Calendário 30 dias': o que postar em cada dia, com o prompt certo. Marque 'Feito'.",
    "6. Aba 'Campanhas do ano': datas para planejar com 3 a 4 semanas de antecedência (Prompt 41).",
    "7. Aba 'Vales e Pacotes': controle os vale-presentes e pacotes vendidos, com validade automática.",
    "",
    "DICAS",
    "• Funciona no Excel e no Google Planilhas (planilhas.google.com > Abrir > Fazer upload).",
    "• Atualize a aba Clientes toda semana. Leva 5 minutos e é o que mais traz dinheiro de volta.",
    "• As datas devem ser digitadas no formato dia/mês/ano, ex.: 15/09/2026.",
])

# ---------------- Ciclos ----------------
ci = wb.create_sheet("Ciclos")
t.titulo(ci, "Ciclo de retorno por serviço", "Em quantos dias a cliente normalmente deve voltar. Ajuste à sua realidade.")
t.cabecalho(ci, 4, ["Serviço", "Ciclo (dias)"])
ciclos = [("Manicure", 10), ("Pedicure", 15), ("Manicure e pedicure", 12), ("Escova", 7), ("Unhas em gel (manutenção)", 25),
          ("Design de sobrancelha", 20), ("Retoque de raiz", 35), ("Coloração", 45), ("Corte feminino", 60),
          ("Corte masculino", 30), ("Barba", 15), ("Mechas / luzes", 90), ("Progressiva / alisamento", 100),
          ("Hidratação / tratamento", 20), ("Cílios (manutenção)", 21), ("Depilação", 30)]
for i in range(25):
    r = 5 + i
    a, b = ci.cell(row=r, column=1), ci.cell(row=r, column=2)
    if i < len(ciclos):
        a.value, b.value = ciclos[i]
    t.entrada(a)
    t.entrada(b)
larguras(ci, [34, 14])

# ---------------- Clientes ----------------
cl = wb.create_sheet("Clientes")
t.titulo(cl, "Clientes: quem chamar hoje", "Preencha nome, serviço e última visita. O status é calculado sozinho.")
cols = ["Nome", "WhatsApp", "Serviço principal", "Última visita", "Ciclo (dias)", "Próximo retorno",
        "Dias de atraso", "Status", "Aniversário", "Mês aniv.", "Observações"]
t.cabecalho(cl, 4, cols)
F, L = 5, 404
dv_serv = DataValidation(type="list", formula1="=Ciclos!$A$5:$A$29", allow_blank=True)
cl.add_data_validation(dv_serv)
dv_serv.add(f"C{F}:C{L}")
for r in range(F, L + 1):
    for c in (1, 2, 3, 11):
        t.entrada(cl.cell(row=r, column=c))
    t.entrada(cl.cell(row=r, column=4), DATA)
    t.entrada(cl.cell(row=r, column=9), 'DD/MM')
    cl.cell(row=r, column=5, value=f'=IF(C{r}="","",IFERROR(VLOOKUP(C{r},Ciclos!$A$5:$B$29,2,FALSE),""))')
    cl.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),"",D{r}+E{r})')
    cl.cell(row=r, column=7, value=f'=IF(F{r}="","",TODAY()-F{r})')
    cl.cell(row=r, column=8, value=(f'=IF(F{r}="","",IF(G{r}<0,"Em dia",IF(G{r}<=15,"Chamar agora",'
                                    f'IF(G{r}<=60,"Sumida","Reativar"))))'))
    cl.cell(row=r, column=10, value=f'=IF(I{r}="","",MONTH(I{r}))')
    t.calculo(cl.cell(row=r, column=5))
    t.calculo(cl.cell(row=r, column=6), DATA)
    t.calculo(cl.cell(row=r, column=7), '0')
    t.calculo(cl.cell(row=r, column=8), negrito=True)
    t.calculo(cl.cell(row=r, column=10))
verde_se(cl, f"H{F}:H{L}", f'$H{F}="Em dia"')
amarelo_se(cl, f"H{F}:H{L}", f'$H{F}="Chamar agora"')
vermelho_se(cl, f"H{F}:H{L}", f'OR($H{F}="Sumida",$H{F}="Reativar")')
# exemplos
ex = [("Ana (exemplo)", "11 99999-0001", "Manicure", -20), ("Carla (exemplo)", "11 99999-0002", "Retoque de raiz", -40),
      ("Bia (exemplo)", "11 99999-0003", "Progressiva / alisamento", -30)]
for i, (n, w, s, d) in enumerate(ex):
    r = F + i
    cl.cell(row=r, column=1, value=n)
    cl.cell(row=r, column=2, value=w)
    cl.cell(row=r, column=3, value=s)
    cl.cell(row=r, column=4, value=f"=TODAY(){d}")
# resumo lateral
t.rotulo(cl, "M4", "RESUMO")
for i, (rot, f) in enumerate([
    ("Clientes cadastradas", f'=COUNTA(A{F}:A{L})'),
    ("Em dia", f'=COUNTIF(H{F}:H{L},"Em dia")'),
    ("Chamar agora", f'=COUNTIF(H{F}:H{L},"Chamar agora")'),
    ("Sumidas", f'=COUNTIF(H{F}:H{L},"Sumida")'),
    ("Reativar (60+ dias)", f'=COUNTIF(H{F}:H{L},"Reativar")'),
    ("Aniversariantes do mês", f'=COUNTIF(J{F}:J{L},MONTH(TODAY()))'),
]):
    cl.cell(row=5 + i, column=13, value=rot)
    t.calculo(cl.cell(row=5 + i, column=14, value=f), negrito=True)
t.nota(cl, "M12", "Apague as 3 linhas de exemplo quando começar.")
t.nota(cl, "M13", "Dica: use Dados > Filtro na linha 4 para ver só 'Chamar agora'.")
larguras(cl, [22, 16, 26, 14, 10, 14, 10, 15, 12, 9, 30, 3, 26, 10])
cl.freeze_panes = "B5"
cl.auto_filter.ref = f"A4:K{L}"

# ---------------- Precificação ----------------
pr = wb.create_sheet("Precificação")
t.titulo(pr, "Precificação: o preço mínimo de cada serviço", "Preço mínimo = (custo do tempo + produto) ÷ (1 − comissão − taxa − margem)")
t.rotulo(pr, "A4", "CUSTOS FIXOS DO MÊS")
fixos = ["Aluguel", "Condomínio / IPTU", "Luz", "Água", "Internet e telefone", "Contador e sistema",
         "Pró-labore (seu salário)", "Salários fixos + encargos", "Limpeza e manutenção", "Marketing", "Outros"]
valores_ex = [2500, 300, 450, 150, 150, 350, 3000, 0, 300, 300, 500]
for i, n in enumerate(fixos):
    r = 5 + i
    pr.cell(row=r, column=1, value=n)
    t.entrada(pr.cell(row=r, column=2, value=valores_ex[i]), BRL)
tf = 5 + len(fixos)
t.rotulo(pr, f"A{tf}", "Total de custos fixos")
t.calculo(pr.cell(row=tf, column=2, value=f"=SUM(B5:B{tf-1})"), BRL, True)

t.rotulo(pr, "D4", "PARÂMETROS")
params = [("Horas produtivas no mês (atendendo de fato)", 250, '0'), ("Comissão padrão da profissional", 0.40, PCT),
          ("Taxa média do cartão / maquininha", 0.03, PCT), ("Margem de lucro desejada", 0.20, PCT)]
for i, (n, v, fmt) in enumerate(params):
    r = 5 + i
    pr.cell(row=r, column=4, value=n)
    c = pr.cell(row=r, column=5, value=v)
    t.entrada(c, fmt)
pr["D10"], pr["D11"] = "Custo por hora", "Custo por minuto"
t.calculo(pr.cell(row=10, column=5, value="=IF(E5=0,0,B16/E5)"), BRL, True)
t.calculo(pr.cell(row=11, column=5, value="=E10/60"), 'R$ #,##0.000', True)
t.nota(pr, "D13", "Horas produtivas: some as horas em que cada profissional está atendendo no mês.")
t.nota(pr, "D14", "Os valores já preenchidos são EXEMPLOS. Troque pelos seus.")

H = 19
t.cabecalho(pr, H, ["Serviço", "Duração (min)", "Custo do produto", "Comissão %", "Preço atual",
                    "Custo do tempo", "PREÇO MÍNIMO", "Diferença", "Lucro por atendimento", "Margem real", "Situação"], 36)
exemplos = [("Escova", 45, 6, 60), ("Corte feminino", 50, 3, 90), ("Coloração", 120, 45, 180), ("Manicure", 40, 4, 35), ("Progressiva", 180, 70, 250)]
for i in range(20):
    r = H + 1 + i
    for c in (1, 2, 3, 4, 5):
        t.entrada(pr.cell(row=r, column=c), {3: BRL, 4: PCT, 5: BRL}.get(c))
    if i < len(exemplos):
        pr.cell(row=r, column=1, value=exemplos[i][0])
        pr.cell(row=r, column=2, value=exemplos[i][1])
        pr.cell(row=r, column=3, value=exemplos[i][2])
        pr.cell(row=r, column=5, value=exemplos[i][3])
    pr.cell(row=r, column=6, value=f'=IF(B{r}="","",B{r}*$E$11)')
    pr.cell(row=r, column=7, value=(f'=IF(B{r}="","",IFERROR((F{r}+C{r})/(1-IF(D{r}="",$E$6,D{r})-$E$7-$E$8),"verifique %"))'))
    pr.cell(row=r, column=8, value=f'=IF(OR(E{r}="",B{r}=""),"",E{r}-G{r})')
    pr.cell(row=r, column=9, value=f'=IF(OR(E{r}="",B{r}=""),"",E{r}*(1-IF(D{r}="",$E$6,D{r})-$E$7)-F{r}-C{r})')
    pr.cell(row=r, column=10, value=f'=IF(OR(E{r}="",E{r}=0,B{r}=""),"",I{r}/E{r})')
    pr.cell(row=r, column=11, value=(f'=IF(OR(E{r}="",B{r}=""),"",IF(E{r}<G{r},"ABAIXO DO MÍNIMO",'
                                     f'IF(E{r}<G{r}*1.15,"No limite","Saudável")))'))
    for c, fmt in [(6, BRL), (7, BRL), (8, BRL), (9, BRL), (10, PCT), (11, None)]:
        t.calculo(pr.cell(row=r, column=c), fmt, c in (7, 11))
vermelho_se(pr, f"K{H+1}:K{H+20}", f'$K{H+1}="ABAIXO DO MÍNIMO"')
amarelo_se(pr, f"K{H+1}:K{H+20}", f'$K{H+1}="No limite"')
verde_se(pr, f"K{H+1}:K{H+20}", f'$K{H+1}="Saudável"')
t.nota(pr, f"A{H+22}", "Comissão % em branco = usa a comissão padrão. Lucro por atendimento já desconta comissão, taxa, tempo e produto.")
larguras(pr, [30, 13, 14, 12, 13, 13, 15, 13, 15, 12, 20])

# ---------------- Metas ----------------
me = wb.create_sheet("Metas")
t.titulo(me, "Metas e ocupação", "Descubra quanto da sua agenda está sendo usada e quanto falta para a meta.")
entradas = [("Número de profissionais atendendo", 2, '0'), ("Horas de trabalho por dia (cada)", 8, '0'),
            ("Dias trabalhados no mês", 22, '0'), ("Atendimentos realizados no mês", 180, '0'),
            ("Duração média de um atendimento (min)", 60, '0'), ("Faturamento do mês (R$)", 16200, BRL),
            ("Meta de faturamento do mês (R$)", 22000, BRL)]
for i, (n, v, fmt) in enumerate(entradas):
    r = 4 + i
    me.cell(row=r, column=1, value=n)
    t.entrada(me.cell(row=r, column=2, value=v), fmt)
calc = [("Horas disponíveis no mês", "=B4*B5*B6", '0'),
        ("Horas atendidas no mês", "=B7*B8/60", '0'),
        ("TAXA DE OCUPAÇÃO", "=IF(B12=0,0,B13/B12)", PCT),
        ("Ticket médio (R$)", "=IF(B7=0,0,B9/B7)", BRL),
        ("Atendimentos necessários para a meta", "=IF(B15=0,0,ROUNDUP(B10/B15,0))", '0'),
        ("Atendimentos a mais por mês", "=MAX(0,B16-B7)", '0'),
        ("Atendimentos a mais por dia", "=IF(B6=0,0,ROUNDUP(B17/B6,1))", '0.0'),
        ("Ocupação necessária para a meta", "=IF(B12=0,0,B16*B8/60/B12)", PCT),
        ("Diagnóstico", '=IF(B19>1,"Meta exige mais que 100% da agenda: suba o ticket (combos, preço) ou a equipe",'
                        'IF(B14<0.6,"Ocupação baixa: foque em retenção e divulgação",IF(B14<0.85,"Ocupação saudável: trabalhe ticket e frequência",'
                        '"Agenda quase cheia: hora de reajustar preço ou contratar")))', None)]
for i, (n, f, fmt) in enumerate(calc):
    r = 12 + i
    me.cell(row=r, column=1, value=n).font = Font(bold=n.isupper() or n == "Diagnóstico")
    t.calculo(me.cell(row=r, column=2, value=f), fmt, True)
me["B20"].alignment = Alignment(wrap_text=True)
me.row_dimensions[20].height = 45

t.rotulo(me, "A23", "ACOMPANHAMENTO MENSAL")
t.cabecalho(me, 24, ["Mês", "Faturamento", "Atendimentos", "Ticket médio", "Clientes novas",
                     "Novas que voltaram em 60 dias", "Taxa de retorno", "Ocupação"], 36)
meses = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
for i, m in enumerate(meses):
    r = 25 + i
    me.cell(row=r, column=1, value=m)
    for c, fmt in [(2, BRL), (3, '0'), (5, '0'), (6, '0'), (8, PCT)]:
        t.entrada(me.cell(row=r, column=c), fmt)
    t.calculo(me.cell(row=r, column=4, value=f'=IF(OR(C{r}="",C{r}=0),"",B{r}/C{r})'), BRL)
    t.calculo(me.cell(row=r, column=7, value=f'=IF(OR(E{r}="",E{r}=0),"",F{r}/E{r})'), PCT)
vermelho_se(me, "G25:G36", 'AND(ISNUMBER($G25),$G25<0.4)')
t.nota(me, "A38", "Taxa de retorno abaixo de 40% (em vermelho): o problema é experiência ou falta de lembrete, não marketing.")
larguras(me, [42, 18, 14, 14, 14, 16, 14, 12])

# ---------------- Calendário ----------------
ca = wb.create_sheet("Calendário 30 dias")
t.titulo(ca, "Calendário de conteúdo: 30 dias", "Um tipo de conteúdo por dia. Gere a semana inteira no domingo com o Prompt 12.")
t.cabecalho(ca, 4, ["Dia", "Semana / foco", "Dia da semana", "Tipo", "Ideia de post", "Prompt", "Feito?"])
plano = {
    "Seg": ("Atração", "Dica rápida de cuidado em casa", "6"),
    "Ter": ("Venda", "Horários livres da semana + link do WhatsApp", "15"),
    "Qua": ("Prova", "Antes e depois do carro-chefe", "7"),
    "Qui": ("Prova", "Depoimento de cliente", "11"),
    "Sex": ("Atração", "Reels de transformação", "9"),
    "Sáb": ("Conexão", "Bastidor: um dia no salão / equipe", "8"),
    "Dom": ("Descanso", "Story leve + gerar os posts da semana", "12"),
}
focos = ["1 · Apresentar o salão", "2 · Carro-chefe", "3 · Prova e avaliações", "4 · Combos e próximo mês", "5 · Revisão"]
dias = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
dvs = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
ca.add_data_validation(dvs)
for d in range(30):
    r = 5 + d
    ds = dias[d % 7]
    tipo, ideia, pr_ = plano[ds]
    for c, v in enumerate([d + 1, focos[d // 7], ds, tipo, ideia, pr_], 1):
        ca.cell(row=r, column=c, value=v)
    t.entrada(ca.cell(row=r, column=7))
    dvs.add(f"G{r}")
verde_se(ca, "A5:G34", '$G5="Sim"')
larguras(ca, [6, 26, 14, 12, 48, 9, 9])

# ---------------- Campanhas ----------------
cp = wb.create_sheet("Campanhas do ano")
t.titulo(cp, "Campanhas do ano", "Comece a divulgar 3 a 4 semanas antes. Use o Prompt 41 para montar cada campanha.")
t.cabecalho(cp, 4, ["Mês", "Data / período", "Oportunidade", "Ideia de campanha", "Começar a divulgar em", "Minha oferta", "Resultado"])
camp = [("Janeiro", "Todo o mês", "Verão e volta das festas", "Pacote de hidratação pós-praia"),
        ("Fevereiro", "Carnaval", "Carnaval", "Tranças, glitter, unhas coloridas, horário estendido"),
        ("Março", "08/03", "Dia da Mulher", "Dia de beleza para presentear, sorteio entre clientes"),
        ("Abril", "Páscoa", "Páscoa", "Mimo de chocolate no atendimento"),
        ("Maio", "2º domingo", "Dia das Mães", "Vale-presente e dia de mãe e filha"),
        ("Junho", "12/06", "Dia dos Namorados", "Combo pronta para o encontro"),
        ("Julho", "Férias", "Férias escolares", "Horários para mães, serviços infantis"),
        ("Agosto", "2º domingo", "Dia dos Pais", "Corte e barba, vale-presente"),
        ("Setembro", "15/09", "Dia do Cliente", "Agradecimento às clientes fiéis"),
        ("Outubro", "Todo o mês", "Outubro Rosa / lista VIP Black Friday", "Ação solidária e abertura da lista VIP"),
        ("Novembro", "Última sexta", "Black Friday", "Combos e vale-presente com bônus (Prompt 42)"),
        ("Dezembro", "Festas", "Natal, Réveillon, formaturas", "Agenda antecipada e pacote das festas")]
for i, row in enumerate(camp):
    r = 5 + i
    for c, v in enumerate(row, 1):
        cp.cell(row=r, column=c, value=v)
    for c in (5, 6, 7):
        t.entrada(cp.cell(row=r, column=c))
larguras(cp, [12, 14, 34, 46, 20, 30, 24])

# ---------------- Vales ----------------
va = wb.create_sheet("Vales e Pacotes")
t.titulo(va, "Vale-presentes e pacotes vendidos", "A validade é calculada sozinha (90 dias). Mude o prazo na célula H2.")
va["G2"], va["H2"] = "Validade (dias):", 90
t.entrada(va["H2"], '0')
t.cabecalho(va, 4, ["Código", "Tipo", "Comprador(a)", "Para quem", "Valor pago", "Valor de uso", "Data da venda",
                    "Válido até", "Já usado?", "Situação"])
dvt = DataValidation(type="list", formula1='"Vale-presente,Pacote mensal,Fidelidade"', allow_blank=True)
dvu = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
va.add_data_validation(dvt)
va.add_data_validation(dvu)
for r in range(5, 105):
    va.cell(row=r, column=1, value=f'=IF(C{r}="","","V"&TEXT(ROW()-4,"000"))')
    t.calculo(va.cell(row=r, column=1))
    for c, fmt in [(2, None), (3, None), (4, None), (5, BRL), (6, BRL), (7, DATA), (9, None)]:
        t.entrada(va.cell(row=r, column=c), fmt)
    va.cell(row=r, column=8, value=f'=IF(G{r}="","",G{r}+$H$2)')
    t.calculo(va.cell(row=r, column=8), DATA)
    va.cell(row=r, column=10, value=f'=IF(G{r}="","",IF(I{r}="Sim","Usado",IF(TODAY()>H{r},"Vencido",IF(H{r}-TODAY()<=15,"Vence em breve","Ativo"))))')
    t.calculo(va.cell(row=r, column=10), negrito=True)
dvt.add("B5:B104")
dvu.add("I5:I104")
amarelo_se(va, "J5:J104", '$J5="Vence em breve"')
vermelho_se(va, "J5:J104", '$J5="Vencido"')
verde_se(va, "J5:J104", '$J5="Ativo"')
va["L4"], va["L5"], va["L6"] = "Total vendido", "Ativos (valor de uso)", "Vencem em breve"
t.calculo(va.cell(row=4, column=13, value="=SUM(E5:E104)"), BRL, True)
t.calculo(va.cell(row=5, column=13, value='=SUMIF(J5:J104,"Ativo",F5:F104)+SUMIF(J5:J104,"Vence em breve",F5:F104)'), BRL, True)
t.calculo(va.cell(row=6, column=13, value='=COUNTIF(J5:J104,"Vence em breve")'), '0', True)
larguras(va, [9, 16, 22, 22, 12, 12, 13, 13, 10, 16, 3, 22, 14])

wb.save(sys.argv[1])
print("ok", sys.argv[1])
