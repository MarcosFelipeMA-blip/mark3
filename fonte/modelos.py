# Gera os modelos editáveis de proposta e contrato em produto/modelos/
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROXO = RGBColor(0x5B, 0x3D, 0xF5)

def base():
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
    return d

def h(d, txt, lvl=1):
    p = d.add_heading(txt, level=lvl)
    for r in p.runs: r.font.color.rgb = ROXO
    return p

def tabela(d, linhas, header=True):
    t = d.add_table(rows=len(linhas), cols=len(linhas[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(linhas):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
            if header and i == 0:
                for r in t.cell(i, j).paragraphs[0].runs: r.bold = True
    d.add_paragraph()

# ---------------- Proposta ----------------
d = base()
t = d.add_paragraph(); r = t.add_run("PROPOSTA COMERCIAL"); r.bold = True; r.font.size = Pt(24); r.font.color.rgb = ROXO
d.add_paragraph("[Seu nome / sua marca] · [Serviço]\nPara: [Empresa do cliente] · A/C [Nome do contato]\nData: [dd/mm/aaaa] · Validade: 7 dias")
h(d, "1. Diagnóstico")
d.add_paragraph("Na nossa conversa em [data], você comentou que [principal problema do cliente]. Hoje, [situação atual — ex.: o perfil está sem publicações regulares desde março e não há padrão visual]. Isso faz com que [consequência — ex.: potenciais clientes não encontrem informações atualizadas e escolham concorrentes].")
h(d, "2. Objetivo")
d.add_paragraph("Nos próximos [3] meses, o objetivo é [objetivo mensurável e realista — ex.: manter presença profissional e constante no Instagram, com calendário mensal e relatório de resultados].")
h(d, "3. Como vamos trabalhar")
for s in ["Semana 1: briefing, guia da marca e calendário do mês.",
          "Semanas 2–4: produção, aprovação em lote e publicação.",
          "Fim do mês: relatório com números, aprendizados e plano do próximo mês."]:
    d.add_paragraph(s, style="List Bullet")
h(d, "4. Opções de investimento")
tabela(d, [["", "Essencial", "Profissional ★", "Completo"],
           ["[Entregável 1]", "[8 posts]", "[12 posts]", "[16 posts]"],
           ["[Entregável 2]", "—", "[Stories 3x/sem]", "[Stories 5x/sem]"],
           ["[Entregável 3]", "—", "[2 reels]", "[4 reels]"],
           ["Relatório mensal", "—", "✓", "✓ + reunião"],
           ["Revisões por peça", "1", "2", "2"],
           ["Investimento mensal", "R$ [   ]", "R$ [   ]", "R$ [   ]"]])
d.add_paragraph("★ Opção recomendada para o seu objetivo.")
h(d, "5. Não está incluso")
d.add_paragraph("Investimento em anúncios pagos, sessões de fotos profissionais, impressão e qualquer item não listado acima. Itens extras podem ser orçados à parte.")
h(d, "6. Condições")
for s in ["Pagamento: mensal e antecipado, via Pix ou boleto, até o dia [5] de cada mês.",
          "Contrato mínimo: [3] meses; depois, renovação mensal com aviso prévio de 30 dias.",
          "Início: até [5] dias úteis após a assinatura do contrato e o primeiro pagamento.",
          "Prazo para aprovação: o cliente aprova o conteúdo em até [2] dias úteis após o envio."]:
    d.add_paragraph(s, style="List Bullet")
h(d, "7. Próximos passos")
d.add_paragraph("1. Escolha o pacote e responda esta proposta.\n2. Envio o contrato para assinatura digital.\n3. Após o primeiro pagamento, agendamos o briefing.")
d.add_paragraph("\nFico à disposição para qualquer dúvida.\n\n[Seu nome]\n[WhatsApp] · [E-mail] · [Link do portfólio]")
d.save("produto/modelos/Modelo-Proposta-Comercial.docx")

# ---------------- Contrato ----------------
d = base()
p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("CONTRATO DE PRESTAÇÃO DE SERVIÇOS DIGITAIS"); r.bold = True; r.font.size = Pt(15)
d.add_paragraph("CONTRATANTE: [Razão social ou nome], inscrito(a) no CNPJ/CPF sob o nº [   ], com endereço em [   ], neste ato representado(a) por [nome], doravante denominado(a) CONTRATANTE.")
d.add_paragraph("CONTRATADO(A): [Seu nome ou razão social], inscrito(a) no CNPJ/CPF sob o nº [   ], com endereço em [   ], doravante denominado(a) CONTRATADO(A).")
d.add_paragraph("As partes acima identificadas celebram o presente contrato, que se regerá pelas cláusulas a seguir.")
clausulas = [
    ("Cláusula 1ª — Do objeto", "O presente contrato tem por objeto a prestação, pelo(a) CONTRATADO(A), dos serviços de [descrição do serviço], conforme o pacote [nome do pacote], que compreende: [lista de entregáveis, quantidades e formatos]."),
    ("Cláusula 2ª — Do que não está incluso", "Não estão incluídos no objeto deste contrato: [ex.: investimento em anúncios, produção de fotos e vídeos presenciais, impressão]. Serviços adicionais serão orçados e contratados à parte."),
    ("Cláusula 3ª — Dos prazos e da aprovação", "O(A) CONTRATADO(A) enviará os materiais para aprovação até [prazo]. O(A) CONTRATANTE terá [2] dias úteis para aprovar ou solicitar ajustes. Na ausência de manifestação nesse prazo, o material será considerado aprovado. Atrasos causados pela falta de envio de informações ou aprovações pelo(a) CONTRATANTE prorrogam os prazos na mesma proporção."),
    ("Cláusula 4ª — Das revisões", "Cada peça inclui até [2] rodadas de revisão. Alterações adicionais ou mudanças de escopo após a aprovação poderão ser cobradas à parte, mediante orçamento prévio."),
    ("Cláusula 5ª — Das obrigações do(a) CONTRATANTE", "Fornecer informações, acessos, logotipos, fotos e demais materiais necessários; responder às solicitações dentro dos prazos; e responsabilizar-se pela veracidade das informações fornecidas sobre seus produtos e serviços."),
    ("Cláusula 6ª — Das obrigações do(a) CONTRATADO(A)", "Executar os serviços com qualidade e dentro dos prazos; manter sigilo sobre informações confidenciais do(a) CONTRATANTE; e não utilizar materiais de terceiros sem a devida autorização ou licença."),
    ("Cláusula 7ª — Do valor e da forma de pagamento", "Pelos serviços, o(a) CONTRATANTE pagará o valor de R$ [   ] ([valor por extenso]) [por mês / pelo projeto], via [Pix/boleto/transferência], [até o dia X de cada mês, antecipadamente / 50% na assinatura e 50% na entrega]."),
    ("Cláusula 8ª — Do atraso no pagamento", "O atraso no pagamento sujeitará o(a) CONTRATANTE a multa de [2]% e juros de [1]% ao mês sobre o valor devido. Após [10] dias de atraso, o(a) CONTRATADO(A) poderá suspender a prestação dos serviços até a regularização."),
    ("Cláusula 9ª — Da vigência e da rescisão", "Este contrato vigora por [3] meses a partir da assinatura, renovando-se automaticamente por períodos mensais. Qualquer das partes poderá rescindi-lo mediante aviso prévio por escrito de [30] dias. Os valores referentes a serviços já executados são devidos."),
    ("Cláusula 10ª — Dos direitos de uso", "Após a quitação integral, o(a) CONTRATANTE poderá utilizar livremente os materiais finais entregues. O(A) CONTRATADO(A) poderá exibir os trabalhos em seu portfólio, salvo manifestação contrária por escrito do(a) CONTRATANTE. Arquivos editáveis [estão / não estão] incluídos."),
    ("Cláusula 11ª — Do uso de ferramentas", "O(A) CONTRATADO(A) poderá utilizar ferramentas digitais, inclusive de inteligência artificial, como apoio à produção, sendo responsável pela revisão, adequação e qualidade final dos materiais entregues."),
    ("Cláusula 12ª — Da proteção de dados", "As partes comprometem-se a tratar eventuais dados pessoais a que tiverem acesso em razão deste contrato em conformidade com a Lei nº 13.709/2018 (LGPD), utilizando-os apenas para a execução dos serviços."),
    ("Cláusula 13ª — Do foro", "Fica eleito o foro da comarca de [cidade/UF] para dirimir quaisquer questões oriundas deste contrato."),
]
for tit, txt in clausulas:
    p = d.add_paragraph(); r = p.add_run(tit); r.bold = True
    d.add_paragraph(txt)
d.add_paragraph("\nE, por estarem de acordo, as partes assinam o presente instrumento, inclusive de forma eletrônica.\n\n[Cidade], [dd] de [mês] de [aaaa].")
d.add_paragraph("\n\n_______________________________\nCONTRATANTE\n\n\n_______________________________\nCONTRATADO(A)")
p = d.add_paragraph(); r = p.add_run("Modelo de referência. Adapte à sua realidade e, sempre que possível, consulte um advogado."); r.italic = True; r.font.size = Pt(9)
d.save("produto/modelos/Modelo-Contrato-Prestacao-Servicos.docx")
print("ok")
