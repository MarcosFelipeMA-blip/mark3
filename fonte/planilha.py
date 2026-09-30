# Gera produto/Planilha-Gestao-do-Freela.xlsx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROXO = "5B3DF5"
HEAD = PatternFill("solid", fgColor=ROXO)
INPUT = PatternFill("solid", fgColor="FFF3BF")
RES = PatternFill("solid", fgColor="E9FBF4")
WHITE_B = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16, color=ROXO)
thin = Side(style="thin", color="D9D6EA")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BRL = '"R$" #,##0.00'
PCT = "0%"

wb = Workbook()

# ---------- Instruções ----------
ws = wb.active
ws.title = "Instruções"
ws["A1"] = "Planilha de Gestão do Freela — Kit Freela Digital com IA"; ws["A1"].font = TITLE
linhas = [
    "Como usar:",
    "1. Precificação: preencha as células AMARELAS com seus custos, pró-labore e horas. O valor hora mínimo é calculado automaticamente.",
    "2. Pacotes: informe as horas estimadas de cada pacote. O preço sugerido considera seu valor hora, margem e impostos/taxas.",
    "3. Clientes (CRM): registre cada potencial cliente, a etapa do funil e a data do próximo contato.",
    "4. Caixa: lance entradas e saídas do mês. O resumo mostra o saldo e o total por categoria.",
    "",
    "Legenda: células amarelas = você preenche · células verdes = resultado automático.",
]
for i, t in enumerate(linhas, start=3):
    ws.cell(row=i, column=1, value=t)
ws["A3"].font = BOLD
ws.column_dimensions["A"].width = 120

# ---------- Precificação ----------
p = wb.create_sheet("Precificação")
p["A1"] = "Precificação — valor hora mínimo"; p["A1"].font = TITLE
p["A3"], p["B3"] = "Custos fixos mensais", "Valor"
for c in ("A3", "B3"):
    p[c].fill, p[c].font = HEAD, WHITE_B
custos = [("Internet", 100), ("Energia (parte do trabalho)", 60), ("Ferramentas de IA", 100),
          ("Canva / design", 35), ("Celular", 50), ("DAS do MEI", 80), ("Outros", 0)]
for i, (n, v) in enumerate(custos, start=4):
    p.cell(row=i, column=1, value=n)
    c = p.cell(row=i, column=2, value=v); c.fill, c.number_format = INPUT, BRL
last = 3 + len(custos)
r = last + 1
p.cell(row=r, column=1, value="Total de custos fixos").font = BOLD
c = p.cell(row=r, column=2, value=f"=SUM(B4:B{last})"); c.fill, c.number_format, c.font = RES, BRL, BOLD
tot_custos = f"B{r}"

r += 2
p.cell(row=r, column=1, value="Parâmetros").font = BOLD
params = [
    ("Pró-labore desejado (R$/mês)", 3000, BRL, "prolabore"),
    ("Horas de trabalho por dia", 6, "0", "hdia"),
    ("Dias de trabalho por mês", 22, "0", "dias"),
    ("% do tempo faturável (produção para clientes)", 0.65, PCT, "fat"),
    ("Margem de lucro desejada", 0.30, PCT, "margem"),
    ("Impostos e taxas sobre o preço", 0.06, PCT, "imp"),
]
ref = {}
for n, v, fmt, key in params:
    r += 1
    p.cell(row=r, column=1, value=n)
    c = p.cell(row=r, column=2, value=v); c.fill, c.number_format = INPUT, fmt
    ref[key] = f"B{r}"

r += 2
p.cell(row=r, column=1, value="Resultados").font = BOLD
r += 1
p.cell(row=r, column=1, value="Horas faturáveis por mês")
c = p.cell(row=r, column=2, value=f"={ref['hdia']}*{ref['dias']}*{ref['fat']}"); c.fill, c.number_format = RES, "0.0"
ref["hfat"] = f"B{r}"
r += 1
p.cell(row=r, column=1, value="VALOR HORA MÍNIMO").font = BOLD
c = p.cell(row=r, column=2, value=f"=IF({ref['hfat']}>0,({tot_custos}+{ref['prolabore']})/{ref['hfat']},0)")
c.fill, c.number_format, c.font = RES, BRL, Font(bold=True, size=13)
ref["vh"] = f"B{r}"
r += 1
p.cell(row=r, column=1, value="Valor hora com margem e impostos")
c = p.cell(row=r, column=2, value=f"={ref['vh']}*(1+{ref['margem']})/(1-{ref['imp']})"); c.fill, c.number_format = RES, BRL
ref["vhfinal"] = f"B{r}"
r += 1
p.cell(row=r, column=1, value="Faturamento necessário no mês")
c = p.cell(row=r, column=2, value=f"={ref['vhfinal']}*{ref['hfat']}"); c.fill, c.number_format = RES, BRL
p.column_dimensions["A"].width = 48; p.column_dimensions["B"].width = 18
for row in p.iter_rows(min_row=3, max_row=r, max_col=2):
    for c in row:
        if c.value is not None: c.border = BOX

# ---------- Pacotes ----------
k = wb.create_sheet("Pacotes")
k["A1"] = "Preço dos pacotes"; k["A1"].font = TITLE
k["A2"] = "Preencha as horas estimadas por mês/projeto. O preço usa os parâmetros da aba Precificação."
hdr = ["Pacote", "Entregáveis (resumo)", "Horas estimadas", "Custo (horas × valor hora)", "Preço sugerido", "Preço que vou cobrar", "Lucro estimado"]
for j, h in enumerate(hdr, start=1):
    c = k.cell(row=4, column=j, value=h); c.fill, c.font, c.alignment = HEAD, WHITE_B, Alignment(wrap_text=True, vertical="center")
pacotes = [("Oferta de entrada", "Diagnóstico + 5 artes", 4), ("Essencial", "8 posts/mês + legendas", 10),
           ("Profissional", "12 posts + stories 3x/sem + 2 reels + relatório", 18), ("Completo", "16 posts + stories 5x/sem + 4 reels + reunião", 28)]
P = "Precificação!"
for i, (n, e, h) in enumerate(pacotes, start=5):
    k.cell(row=i, column=1, value=n).fill = INPUT
    k.cell(row=i, column=2, value=e).fill = INPUT
    c = k.cell(row=i, column=3, value=h); c.fill = INPUT
    c = k.cell(row=i, column=4, value=f"=C{i}*{P}{ref['vh']}"); c.fill, c.number_format = RES, BRL
    c = k.cell(row=i, column=5, value=f"=ROUNDUP(C{i}*{P}{ref['vh']}*(1+{P}{ref['margem']})/(1-{P}{ref['imp']}),-1)-0.1"); c.fill, c.number_format = RES, BRL
    c = k.cell(row=i, column=6, value=f"=E{i}"); c.fill, c.number_format = INPUT, BRL
    c = k.cell(row=i, column=7, value=f"=F{i}*(1-{P}{ref['imp']})-D{i}"); c.fill, c.number_format = RES, BRL
    for j in range(1, 8): k.cell(row=i, column=j).border = BOX
for col, w in zip("ABCDEFG", (22, 46, 14, 20, 16, 18, 16)):
    k.column_dimensions[col].width = w
k.row_dimensions[4].height = 32

# ---------- Clientes (CRM) ----------
cr = wb.create_sheet("Clientes (CRM)")
cr["A1"] = "Clientes e prospecção"; cr["A1"].font = TITLE
etapas = ["Lead", "Abordado", "Respondeu", "Diagnóstico", "Proposta enviada", "Fechado", "Perdido"]
cr["A2"] = "Etapas: " + " → ".join(etapas)
hdr = ["Empresa", "Contato", "Nicho", "Canal", "Etapa", "Último contato", "Próximo contato", "Pacote", "Valor mensal", "Observações", "Atrasado?"]
for j, h in enumerate(hdr, start=1):
    c = cr.cell(row=4, column=j, value=h); c.fill, c.font = HEAD, WHITE_B
dv = DataValidation(type="list", formula1='"' + ",".join(etapas) + '"', allow_blank=True)
dvc = DataValidation(type="list", formula1='"Instagram,WhatsApp,Google Maps,LinkedIn,Indicação,Plataforma,Outro"', allow_blank=True)
cr.add_data_validation(dv); cr.add_data_validation(dvc)
dv.add("E5:E300"); dvc.add("D5:D300")
exemplo = ["Clínica Sorriso (exemplo)", "Dra. Ana", "Odontologia", "Instagram", "Proposta enviada", "2026-10-01", "2026-10-04", "Profissional", 890, "Pediu para retornar na sexta"]
for j, v in enumerate(exemplo, start=1):
    cr.cell(row=5, column=j, value=v)
for i in range(5, 301):
    cr.cell(row=i, column=6).number_format = "DD/MM/YYYY"
    cr.cell(row=i, column=7).number_format = "DD/MM/YYYY"
    cr.cell(row=i, column=9).number_format = BRL
    cr.cell(row=i, column=11, value=f'=IF(AND(G{i}<>"",E{i}<>"Fechado",E{i}<>"Perdido",G{i}<TODAY()),"SIM","")')
from datetime import date
cr["F5"], cr["G5"] = date(2026, 10, 1), date(2026, 10, 4)
cr.conditional_formatting.add("K5:K300", CellIsRule(operator="equal", formula=['"SIM"'], fill=PatternFill("solid", fgColor="FFC9C9")))
r0 = 4
cr["M4"], cr["N4"] = "Resumo do funil", "Qtd."
for c in ("M4", "N4"): cr[c].fill, cr[c].font = HEAD, WHITE_B
for i, e in enumerate(etapas, start=5):
    cr.cell(row=i, column=13, value=e)
    c = cr.cell(row=i, column=14, value=f'=COUNTIF($E$5:$E$300,M{i})'); c.fill = RES
cr.cell(row=13, column=13, value="Receita recorrente (fechados)").font = BOLD
c = cr.cell(row=13, column=14, value='=SUMIF($E$5:$E$300,"Fechado",$I$5:$I$300)'); c.fill, c.number_format = RES, BRL
for col, w in zip("ABCDEFGHIJKLMN", (28, 16, 16, 14, 18, 14, 15, 14, 14, 34, 10, 2, 28, 14)):
    cr.column_dimensions[col].width = w
cr.freeze_panes = "A5"

# ---------- Caixa ----------
cx = wb.create_sheet("Caixa")
cx["A1"] = "Caixa do mês"; cx["A1"].font = TITLE
hdr = ["Data", "Descrição", "Tipo", "Categoria", "Valor"]
for j, h in enumerate(hdr, start=1):
    c = cx.cell(row=4, column=j, value=h); c.fill, c.font = HEAD, WHITE_B
dvt = DataValidation(type="list", formula1='"Entrada,Saída"', allow_blank=True)
cats = ["Cliente recorrente", "Projeto pontual", "Ferramentas", "Impostos", "Marketing", "Equipamento", "Pró-labore", "Outros"]
dvk = DataValidation(type="list", formula1='"' + ",".join(cats) + '"', allow_blank=True)
cx.add_data_validation(dvt); cx.add_data_validation(dvk)
dvt.add("C5:C400"); dvk.add("D5:D400")
ex = [(date(2026, 10, 5), "Clínica Sorriso — mensalidade (exemplo)", "Entrada", "Cliente recorrente", 890),
      (date(2026, 10, 6), "Assinatura ferramenta de IA (exemplo)", "Saída", "Ferramentas", 100)]
for i, row in enumerate(ex, start=5):
    for j, v in enumerate(row, start=1):
        cx.cell(row=i, column=j, value=v)
for i in range(5, 401):
    cx.cell(row=i, column=1).number_format = "DD/MM/YYYY"
    cx.cell(row=i, column=5).number_format = BRL
cx["G4"], cx["H4"] = "Resumo", "Valor"
for c in ("G4", "H4"): cx[c].fill, cx[c].font = HEAD, WHITE_B
cx["G5"], cx["H5"] = "Entradas", '=SUMIF($C$5:$C$400,"Entrada",$E$5:$E$400)'
cx["G6"], cx["H6"] = "Saídas", '=SUMIF($C$5:$C$400,"Saída",$E$5:$E$400)'
cx["G7"], cx["H7"] = "Saldo", "=H5-H6"
cx["G7"].font = BOLD
cx["G9"], cx["H9"] = "Por categoria", "Valor"
for c in ("G9", "H9"): cx[c].fill, cx[c].font = HEAD, WHITE_B
for i, cat in enumerate(cats, start=10):
    cx.cell(row=i, column=7, value=cat)
    cx.cell(row=i, column=8, value=f"=SUMIF($D$5:$D$400,G{i},$E$5:$E$400)")
for i in list(range(5, 8)) + list(range(10, 10 + len(cats))):
    c = cx.cell(row=i, column=8); c.fill, c.number_format = RES, BRL
for col, w in zip("ABCDEFGH", (13, 42, 11, 20, 14, 2, 22, 16)):
    cx.column_dimensions[col].width = w
cx.freeze_panes = "A5"

wb.save("produto/Planilha-Gestao-do-Freela.xlsx")
print("ok")
