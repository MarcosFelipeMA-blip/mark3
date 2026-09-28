"""Gera dist/planilha-meu-primeiro-bebe.xlsx (fórmulas automáticas).

Uso: python3 planilha.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("dist/planilha-meu-primeiro-bebe.xlsx")
MOEDA = 'R$ #,##0.00'
PEACH = PatternFill("solid", fgColor="FDE9DF")
SAGE = PatternFill("solid", fgColor="E6F2E7")
LAV = PatternFill("solid", fgColor="EEE8FA")
SUN = PatternFill("solid", fgColor="FDF4D8")
RED = PatternFill("solid", fgColor="FBE3E0")
THIN = Side(style="thin", color="E0D6CE")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=16, color="7A4B3A")

# (categoria, item, prioridade, quantidade sugerida)
ENXOVAL = [
    ("Roupas", "Body manga curta (RN + P)", "Essencial", 9),
    ("Roupas", "Body manga longa (RN + P)", "Essencial", 9),
    ("Roupas", "Macacão", "Essencial", 6),
    ("Roupas", "Calça", "Essencial", 6),
    ("Roupas", "Meias (pares)", "Essencial", 7),
    ("Roupas", "Touca", "Importante", 2),
    ("Roupas", "Manta / cueiro", "Essencial", 3),
    ("Roupas", "Casaquinho", "Importante", 3),
    ("Roupas", "Fralda de pano (boca)", "Essencial", 10),
    ("Higiene", "Toalha com capuz", "Essencial", 3),
    ("Higiene", "Trocador", "Essencial", 1),
    ("Higiene", "Termômetro digital", "Essencial", 1),
    ("Higiene", "Sabonete líquido neutro (indicado pelo pediatra)", "Essencial", 1),
    ("Higiene", "Algodão / gaze", "Essencial", 1),
    ("Higiene", "Lenços umedecidos sem perfume", "Importante", 4),
    ("Higiene", "Banheira", "Importante", 1),
    ("Higiene", "Cortador de unha infantil", "Importante", 1),
    ("Higiene", "Cesto de roupas", "Importante", 1),
    ("Higiene", "Lixeira com tampa", "Importante", 1),
    ("Sono", "Berço com selo Inmetro", "Essencial", 1),
    ("Sono", "Colchão firme do tamanho do berço", "Essencial", 1),
    ("Sono", "Lençol com elástico", "Essencial", 3),
    ("Sono", "Saco de dormir para bebê", "Importante", 1),
    ("Quarto", "Cômoda", "Importante", 1),
    ("Quarto", "Luz noturna", "Importante", 1),
    ("Quarto", "Cortina / persiana", "Importante", 1),
    ("Quarto", "Poltrona de amamentação", "Pode esperar", 1),
    ("Passeio", "Bebê-conforto com selo Inmetro", "Essencial", 1),
    ("Passeio", "Carrinho", "Importante", 1),
    ("Passeio", "Bolsa de maternidade", "Importante", 1),
    ("Passeio", "Sling / canguru ergonômico", "Importante", 1),
    ("Alimentação", "Absorvente para seios", "Importante", 1),
    ("Alimentação", "Almofada de amamentação", "Pode esperar", 1),
    ("Alimentação", "Cadeirinha de alimentação", "Pode esperar", 1),
    ("Mãe", "Sutiã de amamentação", "Importante", 2),
    ("Mãe", "Absorvente pós-parto (pacotes)", "Essencial", 3),
]
LINHAS_EXTRAS = 30


def cab(ws, row, headers, fill=PEACH):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = BOLD
        cell.fill = fill
        cell.border = BOX
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def larguras(ws, widths):
    for letra, w in widths.items():
        ws.column_dimensions[letra].width = w


def aba_enxoval(wb):
    ws = wb.active
    ws.title = "Enxoval"
    ws["A1"] = "🛒 Planilha do enxoval"
    ws["A1"].font = TITLE
    ws["A2"] = "Preencha apenas as colunas brancas: quantidade, valor unitário e 'Comprado?'. O total é calculado sozinho."
    cab(ws, 4, ["Categoria", "Produto", "Prioridade", "Quantidade", "Valor unitário", "Total", "Comprado?", "Onde / observação"])
    first = 5
    last = first + len(ENXOVAL) + LINHAS_EXTRAS - 1
    for i in range(first, last + 1):
        if i - first < len(ENXOVAL):
            cat, item, prio, qtd = ENXOVAL[i - first]
            ws.cell(row=i, column=1, value=cat)
            ws.cell(row=i, column=2, value=item)
            ws.cell(row=i, column=3, value=prio)
            ws.cell(row=i, column=4, value=qtd)
            ws.cell(row=i, column=7, value="Não")
        ws.cell(row=i, column=5).number_format = MOEDA
        f = ws.cell(row=i, column=6, value=f'=IF(OR(D{i}="",E{i}=""),0,D{i}*E{i})')
        f.number_format = MOEDA
        f.fill = SUN
        for c in range(1, 9):
            ws.cell(row=i, column=c).border = BOX

    dv_prio = DataValidation(type="list", formula1='"Essencial,Importante,Pode esperar,Pense antes"', allow_blank=True)
    dv_sn = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    ws.add_data_validation(dv_prio)
    ws.add_data_validation(dv_sn)
    dv_prio.add(f"C{first}:C{last}")
    dv_sn.add(f"G{first}:G{last}")
    ws.conditional_formatting.add(f"A{first}:H{last}", CellIsRule(operator="equal", formula=['"Sim"'], fill=SAGE))
    ws.freeze_panes = "A5"
    larguras(ws, {"A": 14, "B": 46, "C": 14, "D": 12, "E": 15, "F": 15, "G": 12, "H": 30})
    return first, last


def aba_resumo(wb, first, last):
    ws = wb.create_sheet("Resumo", 0)
    ws["A1"] = "💰 Resumo financeiro — Meu Primeiro Bebê"
    ws["A1"].font = TITLE
    ws["A2"] = "Digite o orçamento disponível na célula amarela. O restante é automático."
    rng_total = f"Enxoval!F{first}:F{last}"
    rng_comp = f"Enxoval!G{first}:G{last}"
    rng_prio = f"Enxoval!C{first}:C{last}"
    linhas = [
        ("TOTAL DO ENXOVAL", f"=SUM({rng_total})", PEACH),
        ("JÁ GASTAMOS", f'=SUMIF({rng_comp},"Sim",{rng_total})', SAGE),
        ("AINDA FALTA", "=B4-B5", LAV),
        ("ORÇAMENTO DISPONÍVEL", None, SUN),
        ("DIFERENÇA (orçamento − total)", "=B7-B4", PEACH),
    ]
    for r, (rot, form, fill) in enumerate(linhas, start=4):
        a = ws.cell(row=r, column=1, value=rot)
        a.font = BOLD
        a.fill = fill
        a.border = BOX
        b = ws.cell(row=r, column=2, value=form)
        b.number_format = MOEDA
        b.border = BOX
        b.font = Font(bold=True, size=13)
        if form is None:
            b.fill = SUN
    ws.conditional_formatting.add("B8", CellIsRule(operator="lessThan", formula=["0"], fill=RED, font=Font(bold=True, color="B03A2E")))
    ws["C8"] = '=IF(B7="","← informe o orçamento",IF(B8<0,"Acima do orçamento: revise itens Pode esperar/Pense antes","Dentro do orçamento 🎉"))'

    ws["A11"] = "Total por prioridade"
    ws["A11"].font = BOLD
    cab(ws, 12, ["Prioridade", "Total", "Já comprado"], fill=LAV)
    for r, p in enumerate(["Essencial", "Importante", "Pode esperar", "Pense antes"], start=13):
        ws.cell(row=r, column=1, value=p).border = BOX
        c = ws.cell(row=r, column=2, value=f'=SUMIF({rng_prio},A{r},{rng_total})')
        c.number_format = MOEDA
        c.border = BOX
        d = ws.cell(row=r, column=3, value=f'=SUMIFS({rng_total},{rng_prio},A{r},{rng_comp},"Sim")')
        d.number_format = MOEDA
        d.border = BOX

    ws["A19"] = "Itens comprados"
    ws["B19"] = f'=COUNTIF({rng_comp},"Sim")&" de "&COUNTA(Enxoval!B{first}:B{last})'
    ws["A19"].font = BOLD
    larguras(ws, {"A": 34, "B": 20, "C": 58})


def aba_fraldas(wb):
    ws = wb.create_sheet("Fraldas")
    ws["A1"] = "🧷 Calculadora de fraldas"
    ws["A1"].font = TITLE
    ws["A2"] = "Estimativas aproximadas — ajuste fraldas/dia e dias no tamanho conforme o seu bebê. Quantidades por pacote variam por marca."
    cab(ws, 4, ["Tamanho", "Fraldas por dia", "Dias no tamanho", "Necessárias", "Recebidas", "Compradas", "Falta", "Fraldas por pacote", "Pacotes a comprar"])
    dados = [("RN", 10, 21, 36), ("P", 9, 45, 46), ("M", 8, 120, 40), ("G", 6, 180, 34), ("XG", 5, 180, 30)]
    for r, (t, dia, dias, pac) in enumerate(dados, start=5):
        vals = [t, dia, dias, f"=B{r}*C{r}", 0, 0, f"=MAX(0,D{r}-E{r}-F{r})", pac, f'=IF(H{r}>0,ROUNDUP(G{r}/H{r},0),"")']
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BOX
            if c in (4, 7, 9):
                cell.fill = SUN
    r = 10
    ws.cell(row=r, column=1, value="TOTAL").font = BOLD
    for c, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (9, "I")):
        ws.cell(row=r, column=c, value=f"=SUM({col}5:{col}9)").font = BOLD
    ws["A12"] = "Dica: compre aos poucos — o bebê pode pular um tamanho. No chá de bebê, peça tamanhos P, M e G."
    larguras(ws, {"A": 12, "B": 15, "C": 16, "D": 13, "E": 12, "F": 12, "G": 10, "H": 18, "I": 18})


def aba_gastos(wb):
    ws = wb.create_sheet("Gastos da gestação")
    ws["A1"] = "📋 Outros gastos (consultas, exames, remédios, transporte…)"
    ws["A1"].font = TITLE
    cab(ws, 3, ["Data", "Descrição", "Categoria", "Valor", "Pago?"])
    dv = DataValidation(type="list", formula1='"Consulta,Exame,Medicamento,Transporte,Maternidade,Outros"', allow_blank=True)
    dv_sn = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv_sn)
    dv.add("C4:C103")
    dv_sn.add("E4:E103")
    for i in range(4, 104):
        ws.cell(row=i, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=i, column=4).number_format = MOEDA
        for c in range(1, 6):
            ws.cell(row=i, column=c).border = BOX
    ws["G3"] = "TOTAL"
    ws["G3"].font = BOLD
    ws["H3"] = "=SUM(D4:D103)"
    ws["H3"].number_format = MOEDA
    ws["G4"] = "Pago"
    ws["H4"] = '=SUMIF(E4:E103,"Sim",D4:D103)'
    ws["H4"].number_format = MOEDA
    ws.freeze_panes = "A4"
    larguras(ws, {"A": 12, "B": 40, "C": 16, "D": 14, "E": 8, "G": 10, "H": 16})


def aba_cha(wb):
    ws = wb.create_sheet("Chá de bebê")
    ws["A1"] = "🎁 Chá de bebê"
    ws["A1"].font = TITLE
    cab(ws, 3, ["Convidado(a)", "Confirmou?", "Tamanho de fralda sugerido", "Presente recebido", "Agradeci?"])
    dv_sn = DataValidation(type="list", formula1='"Sim,Não"', allow_blank=True)
    dv_t = DataValidation(type="list", formula1='"P,M,G,XG"', allow_blank=True)
    ws.add_data_validation(dv_sn)
    ws.add_data_validation(dv_t)
    dv_sn.add("B4:B83")
    dv_sn.add("E4:E83")
    dv_t.add("C4:C83")
    for i in range(4, 84):
        for c in range(1, 6):
            ws.cell(row=i, column=c).border = BOX
    ws["G3"] = "Confirmados"
    ws["H3"] = '=COUNTIF(B4:B83,"Sim")'
    ws["G4"] = "Agradecimentos pendentes"
    ws["H4"] = '=COUNTIFS(D4:D83,"<>",E4:E83,"<>Sim")'
    larguras(ws, {"A": 30, "B": 12, "C": 24, "D": 32, "E": 11, "G": 26, "H": 10})


def main():
    OUT.parent.mkdir(exist_ok=True)
    wb = Workbook()
    first, last = aba_enxoval(wb)
    aba_resumo(wb, first, last)
    aba_fraldas(wb)
    aba_gastos(wb)
    aba_cha(wb)
    wb.active = 0
    wb.save(OUT)
    print("XLSX  ->", OUT)


if __name__ == "__main__":
    main()
